"""v2 工程を run で最後まで通し、読者別 PPTX の構造を確認する（仕様 v2 §7, §13, §14）。"""
from __future__ import annotations

import copy
import re
import zipfile

import pytest
from pptx import Presentation

from caqnal_report.run import Run
from caqnal_report.steps import run_until


@pytest.fixture(scope="module")
def run(cfg, tmp_path_factory):
    c = copy.deepcopy(cfg)
    c.raw["output_dir"] = str(tmp_path_factory.mktemp("out"))
    c.raw.update({"publication": {"min_n": 5}, "aggregate": {"category_missing": "all_answered"},
                  "common9": {"status": "not_conducted"},
                  "models": {"cv": {"group_level": "部局", "folds": 5}, "tree": {"max_depth": 2, "min_leaf": 40},
                             "bootstrap": {"n": 10, "seed": 0}},
                  "readers": [{"id": "全庁", "kind": "all_staff"}, {"id": "総務部長", "kind": "dept_head", "unit": "U003"},
                              {"id": "監査委員事務局長", "kind": "dept_head", "unit": "U016"}]})
    c.raw["org_dictionary"] = "org.yaml"
    c.org_dictionary = {"levels": ["部局"], "units": [{"id": f"U{i:03d}", "name": n, "parent": None, "level": "部局",
                                                      "order": i} for i, n in enumerate(c.org_units, 1)]}
    r = Run.create(c, "test")
    assert run_until(r, "render")
    return r


def _texts(path) -> list[str]:
    prs = Presentation(str(path))
    out = []
    for s in prs.slides:
        out.append("\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                   + "\n".join(c.text for sh in s.shapes if sh.has_table for row in sh.table.rows for c in row.cells))
    return out


def test_decks_written_and_editable(run):
    files = sorted(p.name for p in run.step_dir("render").glob("*.pptx"))
    assert any("詳細付録" in f for f in files) and any("総務部長" in f for f in files)
    main = next(run.step_dir("render").glob("*_全庁レポート.pptx"))
    with zipfile.ZipFile(main) as z:
        charts = [n for n in z.namelist() if re.match(r"ppt/charts/chart\d+\.xml", n)]
        pics = [n for n in z.namelist() if n.startswith("ppt/media/")]
    assert charts and not pics                      # グラフは編集可能なネイティブ chart（画像なし）
    assert run.status == "draft_qa_pending"


def test_hidden_unit_never_shows_values(run):
    app = next(run.step_dir("render").glob("*詳細付録.pptx"))
    texts = _texts(app)
    hidden = [t for t in texts if "監査委員事務局｜" in t]
    assert hidden and all("非表示" in t for t in hidden)
    assert not any("監査委員事務局｜設問詳細" in t for t in texts)


def test_dept_head_sees_only_own_department(run, cfg):
    deck = next(run.step_dir("render").glob("*総務部長*.pptx"))
    body = "\n".join(_texts(deck))
    others = [u for u in cfg.org_units if u not in ("総務部",)]
    assert "総務部" in body
    assert not [u for u in others if u in body]
    assert "アクションプラン記入シート" in body


def test_hidden_own_unit_gets_short_deck_without_values(run):
    deck = next(run.step_dir("render").glob("*監査委員事務局長*.pptx"))
    texts = _texts(deck)
    body = "\n".join(texts)
    assert len(texts) == 4 and "監査委員事務局の結果について" in body
    assert "アクションプラン記入シート" in body and "サマリー" not in body
    assert not re.search(r"\d+\.\d", body)                       # 点数・率は 1 つも無い


def test_font_preset_switches(cfg):
    from caqnal_report.render.v2.base import V2Theme

    c = copy.deepcopy(cfg)
    c.theme["font"]["preset"] = "meiryo"
    assert V2Theme(c).family == "Meiryo UI"
    c.theme["font"]["preset"] = "noto"
    assert V2Theme(c).family == "Noto Sans JP"


def test_appendix_index_pages_match(run):
    app = next(run.step_dir("render").glob("*詳細付録.pptx"))
    texts = _texts(app)
    idx = "\n".join(t for t in texts if "詳細ページ案内" in t)
    m = re.search(r"危機管理部\n部局\n(\d+)～(\d+)", idx)
    assert m, idx[:300]
    first = int(m.group(1))
    assert "A2. 危機管理部｜カテゴリ" in texts[first - 1]


def test_internal_method_deck(run):
    from caqnal_report.render.v2.method import build_method

    path = build_method(run.cfg, run)
    assert path.parent.name == "internal"
    body = "\n".join(_texts(path))
    for word in ("8 工程の流れ", "点数と肯定・中立・否定の計算", "未学習データでの検証", "モデルの採否", "匿名性"):
        assert word in body
    assert "外部配布不可" in body


def test_compose_keeps_all_staff_text_out_of_unit_decks(run, tmp_path):
    import yaml

    from caqnal_report.analysis.compose import compose

    c = copy.deepcopy(run.cfg)
    c.root = tmp_path
    (tmp_path / "interpretation.yaml").write_text(yaml.safe_dump({
        "status": "draft", "insights": {"summary_results": "全庁の見立て"},
        "unit_readers": {"insights": {"summary_results": "部局共通の見立て"}},
        "readers": {"総務部長": {"insights": {"category_scores": "総務部の示唆"}}}}, allow_unicode=True), encoding="utf-8")
    res = compose(c, run.step_dir("publish"), [])
    assert res["全庁"]["insights"]["summary_results"] == "全庁の見立て"
    assert res["総務部長"]["insights"] == {"summary_results": "部局共通の見立て", "category_scores": "総務部の示唆"}
    assert any("下書き" in u for u in res["全庁"]["unresolved"])
