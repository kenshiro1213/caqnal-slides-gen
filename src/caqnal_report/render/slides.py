"""スライド builder（仕様書 §8-3 テンプレート T01〜T15 ＋ 表紙・目次・中表紙・付録）。

各 builder は (slide, ctx, spec) を受け取り、タイトル・リード文・本文を絶対座標で配置する。
scope（全庁 / 部局名）は ctx が持ち、部局別レポートでは同じ builder が部局スコープで動く。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from ..analysis.drivers import ZONE_LABELS
from ..analysis.judgement import SIGNAL_LABELS, SIGNALS
from ..analysis.pipeline import AnalysisResult
from ..analysis.segmentation import QUADRANTS
from ..config import ClientConfig
from . import charts as ch
from .charts import fmt, fmt_pct
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from .pptx_helpers import (add_footnote, add_rect, add_text, add_table, bullets, kpi_card, place_image, set_lead,
                           set_run_font, set_subtitle, set_title)
from .theme import Theme

REGISTRY: dict[str, callable] = {}


def slide_type(name: str):
    def deco(fn):
        REGISTRY[name] = fn
        return fn
    return deco


# --------------------------------------------------------------------------- context
@dataclass
class Ctx:
    cfg: ClientConfig
    res: AnalysisResult
    theme: Theme
    charts_dir: Path
    scope_type: str = "全庁"
    scope: str = "全庁"
    deck: dict = field(default_factory=dict)

    # ---- 基本情報
    @property
    def S(self) -> pd.DataFrame:
        return self.res["scores"]

    @property
    def summary(self) -> dict:
        return self.res.summary

    @property
    def is_all(self) -> bool:
        return self.scope_type == "全庁"

    @property
    def scope_label(self) -> str:
        return f"{self.cfg.client['name']}全体" if self.is_all else self.scope

    @property
    def fy(self) -> str:
        return self.cfg.client.get("fiscal_year_label", self.cfg.fiscal_year)

    @property
    def prev_fy(self) -> str | None:
        return self.summary.get("previous_label") or self.summary.get("previous_fiscal_year")

    @property
    def has_prev(self) -> bool:
        return self.summary.get("previous_fiscal_year") is not None

    @property
    def target(self) -> str:
        return self.cfg.raw["key_driver"]["target"]

    @property
    def target_name(self) -> str:
        return self.cfg.indicator_name(self.target)

    @property
    def org_level(self) -> str:
        return self.cfg.org_level

    @property
    def min_n(self) -> int:
        return self.cfg.analysis["anonymity"]["min_n"]

    def ins(self, stype: str | None = None, scope: str | None = None) -> dict | None:
        stype = stype or self.scope_type
        scope = scope or self.scope
        if stype == "全庁":
            return self.res.insights.get("全庁")
        return (self.res.insights.get(stype) or {}).get(scope)

    # ---- スコア取得
    def rows(self, utype: str, stype: str | None = None, scope: str | None = None) -> pd.DataFrame:
        stype = stype or self.scope_type
        scope = scope or self.scope
        return self.S[(self.S["scope_type"] == stype) & (self.S["scope"] == scope) & (self.S["unit_type"] == utype)]

    def row(self, utype: str, unit: str, stype: str | None = None, scope: str | None = None) -> pd.Series | None:
        d = self.rows(utype, stype, scope)
        d = d[d["unit"] == unit]
        return d.iloc[0] if len(d) else None

    def val(self, utype: str, unit: str, col: str = "mean", stype: str | None = None, scope: str | None = None) -> float:
        r = self.row(utype, unit, stype, scope)
        return float("nan") if r is None else r[col]

    def n_scope(self) -> int:
        r = self.row("指標", self.target)
        return int(r["n"]) if r is not None else 0

    # ---- 設問
    def qtext(self, qid: str, n: int | None = None) -> str:
        t = self.cfg.q(qid)["text"]
        return t if n is None else (t if len(t) <= n else t[: n - 1] + "…")

    def qlabel(self, qid: str, n: int = 26) -> str:
        return f"{qid} {self.qtext(qid, n)}"

    def chart(self, name: str) -> Path:
        return self.charts_dir / f"{name}.png"

    def content(self) -> tuple[float, float, float, float]:
        return self.theme.box("content")

    def footnote_std(self, extra: str = "") -> str:
        s = f"n={self.n_scope():,}。{self.min_n}名未満の組織・属性は非表示（－）。"
        if self.has_prev:
            s += f"前年差は{self.prev_fy}比、※は文言変更等による参考比較。"
        return s + extra


def _frame(slide, ctx: Ctx, spec: dict, title: str, lead: str | None = None, footnote: str | None = None) -> None:
    set_title(slide, ctx.theme, spec.get("title", title))
    lead = spec.get("lead", lead)
    if lead:
        set_lead(slide, ctx.theme, lead)
    add_footnote(slide, ctx.theme, spec.get("footnote", footnote if footnote is not None else ctx.footnote_std()))


def _sig_tuple(v, flag, nd=2) -> tuple[str, int]:
    return (fmt(v, nd, signed=True), int(flag) if not (isinstance(flag, float) and np.isnan(flag)) else 0)


def _sec_no(spec: dict):
    """YAML 1.1 では `no:` キーが False として読まれるため両方を見る。"""
    v = spec.get("no")
    return v if v is not None else spec.get(False)


def _ref_mark(status) -> str:
    return "※" if status == "参考比較" else ""


def short_cat(name: str) -> str:
    """表示用にカテゴリ名の「について」を落とす。"""
    return str(name).replace("について", "") if name else ""


# --------------------------------------------------------------------------- 表紙 / 目次 / 中表紙
@slide_type("cover")
def s_cover(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    c = ctx.cfg.client
    title = spec.get("title") or c.get("report_title", "職員エンゲージメント調査 分析レポート")
    sub = c["name"]
    if spec.get("subtitle_suffix"):
        sub += " " + spec["subtitle_suffix"].format(scope=ctx.scope)
    set_title(slide, th, title)
    set_subtitle(slide, th, f"{sub}\n{ctx.fy}　{c.get('report_date', '')}", size=12)
    add_text(slide, th, 0.47, 0.64, 4.0, 0.4, c.get("honorific", c["name"] + " 御中"), size=12, bold=True)


@slide_type("toc")
def s_toc(slide, ctx: Ctx, spec: dict):
    """テンプレ 2 枚目と同じ: 左に「目次」（CENTER_TITLE）、右の SUBTITLE 枠に章タイトルを 18pt 太字で列挙。"""
    th = ctx.theme
    secs = [s for s in ctx.deck.get("slides", []) if s.get("type") == "section"]
    lines = []
    for s in secs:
        no = _sec_no(s)
        lines.append(f"{no}.  {s['title']}" if no is not None else s["title"])
    size = 18 if len(lines) <= 9 else 14 if len(lines) <= 12 else 12
    ph_title = ph_body = None
    for ph in slide.placeholders:
        t = str(ph.placeholder_format.type)
        if "CENTER_TITLE" in t or ph.placeholder_format.idx == 0:
            ph_title = ph
        elif "SUBTITLE" in t or ph.placeholder_format.idx == 1:
            ph_body = ph
    if ph_title is not None:
        ph_title.left, ph_title.top, ph_title.width, ph_title.height = Inches(0.78), Inches(2.56), Inches(2.73), Inches(0.50)
        tf = ph_title.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = spec.get("title", "目次")
        set_run_font(r, th, 20, True, th.hexv("ink_soft"))
    if ph_body is not None:
        ph_body.left, ph_body.top, ph_body.width, ph_body.height = Inches(3.66), Inches(0.07), Inches(5.84), Inches(5.50)
        tf = ph_body.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.line_spacing = 1.5
            r = p.add_run()
            r.text = line
            set_run_font(r, th, size, True, th.hexv("ink_soft"))


@slide_type("section")
def s_section(slide, ctx: Ctx, spec: dict):
    no = _sec_no(spec)
    t = f"{no}. {spec['title']}" if no is not None else spec["title"]
    set_title(slide, ctx.theme, t)
    if spec.get("subtitle"):
        set_subtitle(slide, ctx.theme, spec["subtitle"])


# --------------------------------------------------------------------------- T01 Executive Summary
@slide_type("exec_onepage")
def s_exec_onepage(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    ins = ctx.ins() or {}
    sen = ins.get("sentences", {})
    _frame(slide, ctx, spec, "Executive Summary", sen.get("fact", ""),
           footnote=ctx.footnote_std("　Key Driver は相関・回帰上の関連であり因果ではない。施策は検討方向性。"))
    x, y, w, h = ctx.content()
    ov = ctx.row("指標", ctx.target)
    # KPI 行
    cw = (w - 0.3) / 4
    dy = ov["diff_yoy"] if ov is not None else np.nan
    yoy_txt = (f"前年差 {fmt(dy, signed=True)}{_ref_mark(ov['yoy_status'])}" if ov is not None and not np.isnan(dy)
               else ("前年データなし" if not ctx.has_prev else "前年差 －"))
    kpi_card(slide, th, x, y, cw, 1.0, ctx.target_name + "（5点満点）", fmt(ov["mean"]) if ov is not None else "－", yoy_txt)
    kpi_card(slide, th, x + cw + 0.1, y, cw, 1.0, "肯定率", fmt_pct(ov["pos"]) if ov is not None else "－",
             f"否定率 {fmt_pct(ov['neg'])}" if ov is not None else "", accent=th.hexv("positive"))
    rr = ctx.summary.get("response_rate")
    kpi_card(slide, th, x + 2 * (cw + 0.1), y, cw, 1.0, "回答者数 / 回答率",
             f"{ctx.n_scope():,}", f"回答率 {fmt(rr, 0)}%" if rr else "回答率 －（対象者数未設定）", accent=th.hexv("ink_soft"))
    lowlab = ctx.cfg.raw["segmentation"]["labels"]["low_low"]
    kpi_card(slide, th, x + 3 * (cw + 0.1), y, cw, 1.0, f"{lowlab}の構成比",
             fmt_pct((ins.get("segment_mix") or {}).get(lowlab)), "エンゲージメント低 × 環境低", accent=th.hexv("negative"))
    # 3 列
    y2 = y + 1.12
    colw = (w - 0.2) / 3
    h2 = 2.15

    def block(bx, title, items, color):
        add_rect(slide, th, bx, y2, colw, 0.26, color)
        add_text(slide, th, bx, y2, colw, 0.26, title, size=9.5, bold=True, color="FFFFFF", anchor="middle")
        add_rect(slide, th, bx, y2 + 0.26, colw, h2 - 0.26, th.hexv("bg"), th.hexv("line"))
        bullets(slide, th, bx + 0.05, y2 + 0.3, colw - 0.1, h2 - 0.35, items, size=8, marker="", spacing=1.05)

    def two(i, text, detail):
        return [[{"text": f"{i + 1}. ", "bold": True}, {"text": text, "size": 8}],
                [{"text": "　　" + detail, "size": 6.8, "color": th.hexv("ink_soft")}]]

    def qdetail(s, main):
        d = f"平均 {fmt(s['mean'])} / {main}"
        if s.get("diff_all") is not None:
            d += f" / 全庁比 {fmt(s['diff_all'], signed=True)}"
        if s.get("diff_yoy") is not None:
            d += f" / 前年差 {fmt(s['diff_yoy'], signed=True)}"
        return d

    st = sum((two(i, s["text"], qdetail(s, f"肯定率 {fmt_pct(s['pos'])}")) for i, s in enumerate(ins.get("top_strengths", [])[:3])), []) or ["該当なし"]
    block(x, "強み", st, th.hexv("positive"))
    iss = sum((two(i, s["text"], qdetail(s, f"否定率 {fmt_pct(s['neg'])}")) for i, s in enumerate(ins.get("top_issues", [])[:3])), []) or ["該当なし"]
    block(x + colw + 0.1, "重点課題候補", iss, th.hexv("negative"))
    kd = sum((two(i, k["text"], f"r={fmt(k['r'])}" + (f" / β={fmt(k['beta'])}" if k.get("beta") is not None else "")
                  + f" / 現状 {fmt(k['score'])} / {k['zone']}:{k['zone_label']}") for i, k in enumerate(ins.get("key_drivers", [])[:3])), []) or ["該当なし"]
    src = "" if ins.get("key_driver_source", "全庁") == ctx.scope_type else "（全庁の結果）"
    block(x + 2 * (colw + 0.1), f"Key Driver{src}", kd, th.hexv("brand_dark"))
    # 下段
    y3 = y2 + h2 + 0.1
    h3 = y + h - y3
    pr = ctx.res["priority_orgs"]
    if ctx.is_all and len(pr):
        focus = pr[pr["band"] == "重点フォロー候補"]["unit"].tolist()
        watch = pr[pr["band"] == "要注視"]["unit"].tolist()
        good = pr[pr["good_practice"]]["unit"].tolist()
        items = [[{"text": "重点フォロー候補: ", "bold": True}, {"text": "・".join(focus) or "なし"}],
                 [{"text": "要注視: ", "bold": True}, {"text": "・".join(watch) or "なし"}],
                 [{"text": "好事例候補: ", "bold": True}, {"text": "・".join(good) or "なし"}]]
        add_text(slide, th, x, y3, 4.4, 0.24, f"重点{ctx.org_level}（§9-1 加点方式）", size=9, bold=True, color=th.hexv("brand_dark"))
        bullets(slide, th, x, y3 + 0.24, 4.4, h3 - 0.24, items, size=7.5, marker="", spacing=1.0)
    else:
        add_text(slide, th, x, y3, 4.4, 0.24, "位置づけ", size=9, bold=True, color=th.hexv("brand_dark"))
        add_text(slide, th, x, y3 + 0.24, 4.4, h3 - 0.24,
                 f"重点フォロー区分: {ins.get('priority_band') or '－'}"
                 f"（加点 {'－' if ins.get('priority_points') is None else ins['priority_points']}）", size=8)
    add_text(slide, th, x + 4.6, y3, 4.5, 0.24, "施策の方向性（検討候補）", size=9, bold=True, color=th.hexv("brand_dark"))
    add_text(slide, th, x + 4.6, y3 + 0.24, 4.5, h3 - 0.24,
             [sen.get("hint", ""), sen.get("action", "")], size=7.5)


# --------------------------------------------------------------------------- 1. 調査概要
@slide_type("survey_overview")
def s_survey_overview(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    c = ctx.cfg.client
    sm = ctx.summary
    _frame(slide, ctx, spec, "調査概要", "調査の基本設計と回答状況。回答率は対象者数が確定次第反映する。", footnote="")
    x, y, w, h = ctx.content()
    n_attr = len(ctx.cfg.qs(("attribute",)))
    n_free = len(ctx.cfg.qs(("freetext",)))
    sc = ctx.cfg.scale("agree5") if "agree5" in ctx.cfg.raw["scales"] else next(iter(ctx.cfg.raw["scales"].values()))
    labels = "／".join(sc["labels"].keys())
    rr = sm.get("response_rate")
    rows = [["項目", "内容"],
            ["調査名", c.get("survey_title", "職員エンゲージメント調査")],
            ["実施主体", f"{c['name']}（分析: 株式会社CAQNAL）"],
            ["実施時期", c.get("survey_period", "－")],
            ["回答者数", f"{sm['respondents']:,} 名" + (f"（{ctx.prev_fy}: {sm['respondents_prev']:,} 名）" if sm.get("respondents_prev") else "")],
            ["回答率", f"{fmt(rr, 1)}%" if rr else "－（対象者数未設定）"],
            ["設問数", f"尺度設問 {sm['question_count']} 問 ／ 属性 {n_attr} 問 ／ 自由記述 {n_free} 問"],
            ["回答尺度", f"5段階（{labels}）。肯定=上位2、否定=下位2"],
            ["結果指標", f"{ctx.target_name} = " + "・".join(ctx.qtext(q, 18) for q in ctx.cfg.indicator_questions(ctx.target))],
            ["経年比較", (f"{ctx.prev_fy} と比較。比較可 {sm['yoy']['comparable']} 問 ／ 参考比較 {sm['yoy']['reference']} 問"
                          if ctx.has_prev else "前年度データなし")],
            ["匿名性", f"{ctx.min_n} 名未満の組織・属性、{sm['anonymity']['min_comments']} 件未満の自由記述は非表示"]]
    add_table(slide, th, x, y, 5.6, rows, col_widths=[1.2, 4.4], row_h=0.36, size=8, aligns=["left", "left"])
    steps = ["STEP1 調査結果の把握（What）", "STEP2 組織・属性差の把握（Where）", "STEP3 職員タイプの類型化（Who）",
             "STEP4 Key Driver 分析（Why）", "STEP5 管理職WSでの現場確認（So What）", "STEP6 施策化（Now What）"]
    add_text(slide, th, x + 5.9, y, 3.2, 0.3, "分析の流れ", size=10, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 5.9, y + 0.35, 3.2, h - 0.4, steps, size=8.5, spacing=1.35)


@slide_type("respondent_attributes")
def s_respondent_attributes(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    _frame(slide, ctx, spec, "回答者の属性", f"回答者 {ctx.summary['respondents']:,} 名の属性構成。組織別は{ctx.org_level}単位。",
           footnote=f"{ctx.min_n}名未満の組織は個別スコアを非表示とする（{ '・'.join(o['name'] for o in ctx.summary['organizations'] if o['suppressed']) or 'なし'}）。")
    x, y, w, h = ctx.content()
    panels = []
    for key, spec_a in ctx.cfg.attributes.items():
        counts = ctx.summary["attributes"][key]["counts"]
        order = list(spec_a.get("order", [])) + [k for k in counts if k not in spec_a.get("order", [])]
        short = spec_a.get("short", {})
        panels.append((key, [short.get(o, o) for o in order], [counts.get(o, 0) for o in order]))
    orgs = ctx.summary["organizations"]
    panels.append((ctx.org_level, [o["name"] for o in orgs], [o["n"] for o in orgs]))
    org_panel = panels.pop()
    ncol = 3
    lw = 5.9
    cw, chh = (lw - 0.2) / ncol, (h - 0.1) / 2
    for i, (title, labels, counts) in enumerate(panels[:6]):
        cx = x + (i % ncol) * (cw + 0.1)
        cy = y + (i // ncol) * (chh + 0.1)
        add_text(slide, th, cx, cy, cw, 0.25, title, size=9, bold=True, color=th.hexv("brand_dark"))
        p = ch.bar_counts(th, ctx.chart(f"attr_{i}"), labels, counts, w=cw, h=chh - 0.3, font=7, label_chars=9, pct=True)
        place_image(slide, th, p, cx, cy + 0.27, cw, chh - 0.3, align="left")
    title, labels, counts = org_panel
    ox = x + lw + 0.1
    add_text(slide, th, ox, y, w - lw - 0.1, 0.25, title, size=9, bold=True, color=th.hexv("brand_dark"))
    p = ch.bar_counts(th, ctx.chart("attr_org"), labels, counts, w=w - lw - 0.1, h=h - 0.3, font=7, label_chars=12, pct=True)
    place_image(slide, th, p, ox, y + 0.27, w - lw - 0.1, h - 0.3, align="left")


@slide_type("reading_guide")
def s_reading_guide(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    _frame(slide, ctx, spec, "結果の読み方", "本レポートで使う指標と表示ルール。全ページ共通。", footnote="")
    x, y, w, h = ctx.content()
    left = [
        [{"text": "平均（スコア）　", "bold": True}, {"text": "5=そう思う … 1=そう思わない として算出した平均。逆転項目は反転して集計"}],
        [{"text": "肯定率／否定率　", "bold": True}, {"text": "上位2選択肢／下位2選択肢の回答割合（%）。「どちらともいえない」は中立"}],
        [{"text": "カテゴリ・軸・指標　", "bold": True}, {"text": "構成設問の回答者平均。肯定率は構成設問の全回答をプール"}],
        [{"text": "全庁差　", "bold": True}, {"text": f"{ctx.org_level}・属性の平均 − 全庁平均。±{ctx.cfg.analysis['judgement']['org_gap_abs_at_least']} 以上を特徴とみなす"}],
        [{"text": "前年差　", "bold": True}, {"text": f"当年平均 − 前年平均。±{ctx.cfg.analysis['judgement']['improved_at_least']} 以上を改善／悪化。文言変更等は（参考）"}],
        [{"text": "有意差　", "bold": True}, {"text": "Welch の t 検定（α=0.05）。表では青＝有意に高い、桃＝有意に低い。n が小さい場合は検出力不足に注意"}],
        [{"text": "Key Driver　", "bold": True}, {"text": f"{ctx.target_name}との相関・標準化回帰係数。「関連」であり因果ではない"}],
    ]
    bullets(slide, th, x, y, 5.3, h, left, size=8, spacing=1.15)
    lab = ctx.cfg.raw["segmentation"]["labels"]
    right = [
        [{"text": "非表示　", "bold": True}, {"text": f"{ctx.min_n}名未満の組織・属性は「－」。欠損（未回答）とは区別"}],
        [{"text": "色の意味　", "bold": True}, {"text": "青＝高評価・改善・ポジティブ、桃＝低評価・悪化・注意、グレー＝中立。全ページ固定"}],
        [{"text": "4類型　", "bold": True}, {"text": f"{ctx.target_name}×{ctx.cfg.raw['segmentation'].get('y_label', ctx.cfg.raw['segmentation']['y'])}を全庁平均で高低に分け、"
                                              f"{lab['high_high']}／{lab['high_low']}／{lab['low_high']}／{lab['low_low']} に分類"}],
        [{"text": "重点課題候補　", "bold": True}, {"text": "低スコア・悪化・組織偏在・Key Driver・自由記述・ばらつき の6観点のうち3つ以上が該当"}],
        [{"text": "示唆文　", "bold": True}, {"text": "数値から規則的に生成した初期仮説。最終的な解釈・施策は CAQNAL と自治体の協議で確定"}],
    ]
    bullets(slide, th, x + 5.5, y, 3.6, h, right, size=8, spacing=1.15)


# --------------------------------------------------------------------------- T02 総合 KPI
@slide_type("overall_kpi")
def s_overall_kpi(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    ov = ctx.row("指標", ctx.target)
    ins = ctx.ins() or {}
    _frame(slide, ctx, spec, f"{ctx.target_name}", (ins.get("sentences") or {}).get("fact", ""))
    x, y, w, h = ctx.content()
    cw = 2.6
    kpi_card(slide, th, x, y, cw, 1.15, f"{ctx.target_name}（平均）", fmt(ov["mean"]),
             (f"前年 {fmt(ov['prev_mean'])} → 前年差 {fmt(ov['diff_yoy'], signed=True)}{_ref_mark(ov['yoy_status'])}"
              + ("（有意）" if ov["sig_yoy"] else "")) if not np.isnan(ov["diff_yoy"]) else "前年差 －")
    kpi_card(slide, th, x, y + 1.25, cw, 1.0, "肯定率", fmt_pct(ov["pos"]),
             f"前年 {fmt_pct(ov['prev_pos'])}" if not np.isnan(ov["prev_pos"]) else "", accent=th.hexv("positive"), value_size=22)
    kpi_card(slide, th, x, y + 2.35, cw, 1.0, "否定率", fmt_pct(ov["neg"]), f"中立 {fmt_pct(ov['neu'])}",
             accent=th.hexv("negative"), value_size=22)
    kpi_card(slide, th, x, y + 3.45, cw, 0.9, "0-100 換算", fmt(ov["norm"], 1), "（平均−1）÷4×100", accent=th.hexv("ink_soft"), value_size=18)
    # 分布
    qids = ctx.cfg.indicator_questions(ctx.target)
    dist = ctx.res["dist"]
    dist = dist[(dist["scope_type"] == ctx.scope_type) & (dist["scope"] == ctx.scope)].set_index("question_id")
    sc = ctx.cfg.scale(ctx.cfg.q(qids[0]).get("scale") or "agree5")
    order = sorted(sc["labels"].items(), key=lambda kv: -kv[1])
    legend = [k for k, _ in order]
    labels, dists = [], []
    for q in qids:
        if q in dist.index:
            labels.append(ctx.qtext(q))
            dists.append([dist.loc[q, f"s{int(v)}"] for _, v in order])
    px = x + cw + 0.2
    pw = w - cw - 0.2
    add_text(slide, th, px, y, pw, 0.25, "構成設問の回答分布（%）", size=9, bold=True, color=th.hexv("brand_dark"))
    if labels:
        p = ch.dist_stacked(th, ctx.chart("overall_dist"), labels, dists, legend, w=pw, h=2.1, font=7.5)
        place_image(slide, th, p, px, y + 0.27, pw, 2.1, align="left")
    rows = [["構成設問", "平均", ("前年" if ctx.has_prev else "－"), "前年差", "肯定率", "否定率"]]
    fills, colors = {}, {}
    for i, q in enumerate(qids, start=1):
        r = ctx.row("設問", q)
        if r is None:
            continue
        rows.append([ctx.qtext(q), fmt(r["mean"]), fmt(r["prev_mean"]), fmt(r["diff_yoy"], signed=True) + _ref_mark(r["yoy_status"]),
                     fmt_pct(r["pos"]), fmt_pct(r["neg"])])
        if r["sig_yoy"]:
            fills[(i, 3)] = th.hexv("positive_soft") if r["sig_yoy"] > 0 else th.hexv("negative_soft")
    rows.append([f"{ctx.target_name}（3問平均）", fmt(ov["mean"]), fmt(ov["prev_mean"]),
                 fmt(ov["diff_yoy"], signed=True) + _ref_mark(ov["yoy_status"]), fmt_pct(ov["pos"]), fmt_pct(ov["neg"])])
    add_table(slide, th, px, y + 2.5, pw, rows, col_widths=[3.4, 0.7, 0.7, 0.9, 0.7, 0.7], row_h=0.3, size=8,
              fills=fills, bold_cells={(len(rows) - 1, c) for c in range(6)})


# --------------------------------------------------------------------------- T03 カテゴリー結果
@slide_type("category_result")
def s_category_result(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    utype = spec.get("unit_type", "カテゴリ")
    units = ctx.summary["categories"] if utype == "カテゴリ" else ctx.summary["axes"]
    rows = ctx.rows(utype).set_index("unit").reindex(units)
    ref = ctx.val("指標", ctx.target, "mean", "全庁", "全庁") if ctx.is_all else None
    ref_all = None if ctx.is_all else [ctx.val(utype, u, "mean", "全庁", "全庁") for u in units]
    top = rows["mean"].idxmax() if rows["mean"].notna().any() else "－"
    bot = rows["mean"].idxmin() if rows["mean"].notna().any() else "－"
    lead = f"{ctx.scope_label}の{utype}別平均。最も高いのは「{top}」（{fmt(rows['mean'].max())}）、最も低いのは「{bot}」（{fmt(rows['mean'].min())}）。"
    if ctx.has_prev and rows["diff_yoy"].notna().any():
        up, dn = rows["diff_yoy"].idxmax(), rows["diff_yoy"].idxmin()
        lead += f"前年差は「{up}」{fmt(rows.loc[up, 'diff_yoy'], signed=True)} が最大、「{dn}」{fmt(rows.loc[dn, 'diff_yoy'], signed=True)} が最小。"
    _frame(slide, ctx, spec, f"{utype}別結果", lead)
    x, y, w, h = ctx.content()
    labels = [short_cat(u) for u in units]
    diffs = rows["diff_yoy"].tolist() if ctx.has_prev else (rows["diff_all"].tolist() if not ctx.is_all else None)
    sig = rows["sig_yoy"].tolist() if ctx.has_prev else (rows["sig_all"].tolist() if not ctx.is_all else None)
    p = ch.hbar_scores(th, ctx.chart(f"cat_{utype}"), labels, rows["mean"].tolist(),
                       prev=rows["prev_mean"].tolist() if ctx.has_prev else ref_all, diffs=diffs,
                       diff_title="前年差" if ctx.has_prev else "全庁差", ref=ref, ref_label=ctx.target_name,
                       w=6.0, h=h, sig=sig, font=8 if len(units) <= 8 else 7.5, label_chars=16)
    place_image(slide, th, p, x, y, 6.1, h, align="left")
    # 右: 表
    trows = [[utype, "平均", "肯定率", "否定率", "前年差" if ctx.has_prev else "全庁差"]]
    fills = {}
    for i, u in enumerate(units, start=1):
        r = rows.loc[u]
        d, s = (r["diff_yoy"], r["sig_yoy"]) if ctx.has_prev else (r["diff_all"], r["sig_all"])
        trows.append([u, fmt(r["mean"]), fmt_pct(r["pos"]), fmt_pct(r["neg"]), fmt(d, signed=True) + (_ref_mark(r["yoy_status"]) if ctx.has_prev else "")])
        if s:
            fills[(i, 4)] = th.hexv("positive_soft") if s > 0 else th.hexv("negative_soft")
    rh = min(0.3, (h - 0.1) / len(trows))
    add_table(slide, th, x + 6.2, y, 2.9, trows, col_widths=[1.5, 0.5, 0.55, 0.55, 0.7], row_h=rh, size=7, fills=fills)


# --------------------------------------------------------------------------- T04 / T05 ランキング
def _ranking_rows(ctx: Ctx, kind: str, n: int = 10) -> pd.DataFrame:
    rk = ctx.res["rankings"]
    d = rk[(rk["kind"] == kind) & (rk["scope_type"] == ctx.scope_type) & (rk["scope"] == ctx.scope)]
    return d.sort_values("rank").head(n)


def _ranking_table(ctx: Ctx, path: Path, d: pd.DataFrame, w: float, with_dist: bool = True) -> Path:
    dist = ctx.res["dist"]
    dist = dist[(dist["scope_type"] == ctx.scope_type) & (dist["scope"] == ctx.scope)].set_index("question_id")
    narrow = w < 6.5
    if narrow:
        cols = [{"id": "rank", "title": "順位", "width": 0.4, "kind": "num"},
                {"id": "q", "title": "設問", "width": w - 0.4 - 0.55 - 0.8, "kind": "text", "chars": 24, "lines": 2},
                {"id": "mean", "title": "平均", "width": 0.55, "kind": "num"},
                {"id": "yoy", "title": "前年差" if ctx.has_prev else "全庁差", "width": 0.8, "kind": "sig"}]
    else:
        cols = [{"id": "rank", "title": "順位", "width": 0.4, "kind": "num"},
                {"id": "cat", "title": "カテゴリ", "width": 1.35, "kind": "text", "chars": 11},
                {"id": "q", "title": "設問", "width": w - 0.4 - 1.35 - 0.6 * 3 - 0.8 - (1.3 if with_dist else 0), "kind": "text", "chars": 34},
                {"id": "mean", "title": "平均", "width": 0.6, "kind": "num"},
                {"id": "pos", "title": "肯定率", "width": 0.6, "kind": "pct"},
                {"id": "neg", "title": "否定率", "width": 0.6, "kind": "pct"},
                {"id": "yoy", "title": "前年差" if ctx.has_prev else "全庁差", "width": 0.8, "kind": "sig"}]
        if with_dist:
            cols.append({"id": "bar", "title": "肯定/中立/否定", "width": 1.3, "kind": "bar3"})
    rows = []
    for _, r in d.iterrows():
        q = r["question_id"]
        sc = ctx.row("設問", q)
        dv, sg = (r["diff_yoy"], sc["sig_yoy"]) if ctx.has_prev else (r["diff_all"], sc["sig_all"])
        cells = {"rank": int(r["rank"]), "cat": short_cat(ctx.cfg.q(q).get("category", "")), "q": f"{q} {ctx.qtext(q)}",
                 "mean": r["mean"], "pos": r["pos"], "neg": r["neg"],
                 "yoy": (fmt(dv, signed=True) + (_ref_mark(r["yoy_status"]) if ctx.has_prev else ""), sg)}
        if with_dist:
            cells["bar"] = (sc["pos"], sc["neu"], sc["neg"])
        rows.append({"cells": cells})
    return ch.render_table(ctx.theme, path, cols, rows, row_h=0.38 if narrow else 0.3, header_h=0.4, font=7.2 if narrow else 7.5)


@slide_type("question_ranking")
def s_question_ranking(slide, ctx: Ctx, spec: dict):
    kind = spec.get("kind", "高評価")
    d = _ranking_rows(ctx, kind)
    cats = d["question_id"].map(lambda q: ctx.cfg.q(q).get("category")).value_counts()
    common = "・".join(cats.index[:3])
    lead = (f"{ctx.scope_label}の設問別平均 {'上位' if kind == '高評価' else '下位'}{len(d)}。"
            f"共通するテーマは「{common}」。" if len(d) else "")
    _frame(slide, ctx, spec, f"設問別ランキング（{kind}{'TOP' if kind == '高評価' else 'BOTTOM'}{len(d)}）", lead)
    x, y, w, h = ctx.content()
    p = _ranking_table(ctx, ctx.chart(f"rank_{kind}"), d, w)
    place_image(slide, ctx.theme, p, x, y, w, h, align="left")


@slide_type("trend_ranking")
def s_trend_ranking(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    if not ctx.has_prev:
        _frame(slide, ctx, spec, "改善・悪化項目", "前年度データが無いため経年比較は行っていない。")
        return
    up = _ranking_rows(ctx, "改善")
    dn = _ranking_rows(ctx, "悪化")
    j = ctx.cfg.analysis["judgement"]
    n_up = int((up["diff_yoy"] >= j["improved_at_least"]).sum())
    n_dn = int((dn["diff_yoy"] <= j["worsened_at_most"]).sum())
    lead = (f"{ctx.prev_fy}→{ctx.fy} の前年差。改善（{j['improved_at_least']:+.2f}以上）{n_up} 問、"
            f"悪化（{j['worsened_at_most']:+.2f}以下）{n_dn} 問（比較可能な設問のみ）。")
    _frame(slide, ctx, spec, "改善・悪化項目（前年差ランキング）", lead)
    x, y, w, h = ctx.content()
    half = (w - 0.2) / 2
    for i, (title, d, color) in enumerate((("改善 TOP10", up, th.hexv("positive")), ("悪化 BOTTOM10", dn, th.hexv("negative")))):
        bx = x + i * (half + 0.2)
        add_text(slide, th, bx, y, half, 0.25, title, size=9, bold=True, color=color)
        p = _ranking_table(ctx, ctx.chart(f"trend_{i}"), d, half, with_dist=False)
        place_image(slide, th, p, bx, y + 0.27, half, h - 0.3, align="left")


@slide_type("quadrant_yoy")
def s_quadrant_yoy(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    if not ctx.has_prev:
        _frame(slide, ctx, spec, "当年スコア × 前年差", "前年度データが無いため作成していない。")
        return
    jd = ctx.res["judgement"]
    d = jd.dropna(subset=["diff_yoy"])
    labels = ctx.cfg.analysis["quadrant"]["labels"]
    counts = d["quadrant"].value_counts()
    lead = ("設問ごとの当年平均（横）と前年差（縦）。右上=維持、左上=改善途上、右下=要注視、左下=最優先。"
            f"最優先 {counts.get(labels['low_down'], 0)} 問／要注視 {counts.get(labels['high_down'], 0)} 問。")
    _frame(slide, ctx, spec, "当年スコア × 前年差（4象限）", lead)
    x, y, w, h = ctx.content()
    med = float(d["mean"].median())
    cols = [th.c("negative") if q == labels["low_down"] else th.c("brand") if q == labels["high_down"]
            else th.c("positive") if q == labels["high_up"] else th.c("ink_soft") for q in d["quadrant"]]
    p = ch.scatter_labeled(th, ctx.chart("quad_yoy"), d["mean"].tolist(), d["diff_yoy"].tolist(), d["question_id"].tolist(),
                           xline=med, yline=0, xlabel="当年平均", ylabel="前年差", w=6.0, h=h, colors=cols,
                           quadrant_text={"tl": labels["low_up"], "tr": labels["high_up"], "bl": labels["low_down"], "br": labels["high_down"]})
    place_image(slide, th, p, x, y, 6.1, h, align="left")
    pri = d[d["quadrant"] == labels["low_down"]].sort_values("diff_yoy").head(9)
    items = [[{"text": f"{r['question_id']} ", "bold": True}, {"text": f"{r['text']}  {fmt(r['mean'])} / {fmt(r['diff_yoy'], signed=True)}", "size": 7.5}]
             for _, r in pri.iterrows()]
    add_text(slide, th, x + 6.2, y, 2.9, 0.25, f"{labels['low_down']}（低スコア × 悪化）", size=9, bold=True, color=th.hexv("negative"))
    bullets(slide, th, x + 6.2, y + 0.27, 2.9, h - 0.3, items or ["該当なし"], size=7.5, marker="", spacing=1.05)


@slide_type("dispersion")
def s_dispersion(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    d = ctx.rows("設問").dropna(subset=["mean", "sd"])
    d = d[d["unit"].isin(ctx.cfg.likert_ids())]
    thr = float(d["sd"].quantile(1 - ctx.cfg.analysis["judgement"]["high_sd_top_pct"] / 100))
    hi = d[d["sd"] >= thr].sort_values("sd", ascending=False)
    lead = (f"設問ごとの平均（横）と標準偏差（縦）。SD が上位{ctx.cfg.analysis['judgement']['high_sd_top_pct']}%"
            f"（{fmt(thr)} 以上）の {len(hi)} 問は職員間で認識が分かれており、組織内マネジメントの論点になりうる。")
    _frame(slide, ctx, spec, "ばらつき（平均 × 標準偏差）", lead)
    x, y, w, h = ctx.content()
    cols = [th.c("negative") if s >= thr else th.c("positive") for s in d["sd"]]
    p = ch.scatter_labeled(th, ctx.chart("dispersion"), d["mean"].tolist(), d["sd"].tolist(), d["unit"].tolist(),
                           xline=float(d["mean"].median()), yline=thr, xlabel="平均", ylabel="標準偏差", w=6.0, h=h, colors=cols,
                           quadrant_text={"tl": "低評価 × 意見が分かれる", "tr": "高評価 × 意見が分かれる"})
    place_image(slide, th, p, x, y, 6.1, h, align="left")
    items = [[{"text": f"{r['unit']} ", "bold": True}, {"text": f"{ctx.qtext(r['unit'])}  SD {fmt(r['sd'])} / 平均 {fmt(r['mean'])}", "size": 7.5}]
             for _, r in hi.head(10).iterrows()]
    add_text(slide, th, x + 6.2, y, 2.9, 0.25, "ばらつきが大きい設問", size=9, bold=True, color=th.hexv("negative"))
    bullets(slide, th, x + 6.2, y + 0.27, 2.9, h - 0.3, items or ["該当なし"], size=7.5, marker="", spacing=1.05)


# --------------------------------------------------------------------------- T08 4類型
def _seg_rows(ctx: Ctx, stype: str, scope: str) -> pd.DataFrame:
    sg = ctx.res["segments"]
    return sg[(sg["scope_type"] == stype) & (sg["scope"] == scope)].set_index("segment").reindex(QUADRANTS)


def _quadrant_diagram(slide, ctx: Ctx, x, y, w, h, seg: pd.DataFrame, compare: pd.DataFrame | None, compare_label: str):
    th = ctx.theme
    segc = ctx.cfg.raw["segmentation"]
    lab = segc["labels"]
    sm = ctx.summary["segmentation"]
    ax_w, ax_h = 0.35, 0.3
    gx, gy, gw, gh = x + ax_w, y, w - ax_w, h - ax_h
    cw, chh = gw / 2, gh / 2
    pos = {"high_high": (1, 0), "low_high": (0, 0), "high_low": (1, 1), "low_low": (0, 1)}
    for k, (cx, cy) in pos.items():
        share = seg.loc[k, "share"] if k in seg.index else np.nan
        cmp = compare.loc[k, "share"] if compare is not None and k in compare.index else np.nan
        col = th.seg_color(k)
        add_rect(slide, th, gx + cx * cw + 0.03, gy + cy * chh + 0.03, cw - 0.06, chh - 0.06, col)
        sub = f"{compare_label} {fmt_pct(cmp)}（{fmt(share - cmp, 1, signed=True)}pt）" if not np.isnan(cmp) else ""
        add_text(slide, th, gx + cx * cw + 0.08, gy + cy * chh + 0.08, cw - 0.16, chh - 0.16,
                 [[{"text": lab[k], "bold": True, "size": 10, "color": "FFFFFF"}],
                  [{"text": fmt_pct(share), "bold": True, "size": 20, "color": "FFFFFF"}],
                  [{"text": sub, "size": 7.5, "color": "FFFFFF"}]], anchor="middle", align="center")
    add_text(slide, th, x, gy, ax_w, gh, f"{segc.get('y_label', segc['y'])}\n高 ↑\n\n\n\n低 ↓", size=7, color=th.hexv("ink_soft"),
             anchor="middle", align="center")
    add_text(slide, th, gx, gy + gh, gw, ax_h, f"← 低　　{segc.get('x_label', ctx.target_name)}（しきい値 {fmt(sm['threshold_x'])}）　　高 →",
             size=7, color=th.hexv("ink_soft"), align="center")


@slide_type("segment_overview")
def s_segment_overview(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    seg = _seg_rows(ctx, ctx.scope_type, ctx.scope)
    lab = ctx.cfg.raw["segmentation"]["labels"]
    prev = seg["prev_share"] if ctx.has_prev else None
    lead = f"{ctx.scope_label}の職員を{ctx.target_name}×{ctx.cfg.raw['segmentation'].get('y_label', ctx.cfg.raw['segmentation']['y'])}の全庁平均で4類型に分類。"
    lead += f"{lab['high_high']} {fmt_pct(seg.loc['high_high', 'share'])}、{lab['low_low']} {fmt_pct(seg.loc['low_low', 'share'])}。"
    if prev is not None and not np.isnan(seg.loc["low_low", "diff"]):
        lead += f"{lab['low_low']}は前年比 {fmt(seg.loc['low_low', 'diff'], 1, signed=True)}pt。"
    _frame(slide, ctx, spec, "職員タイプ（4類型）の構成", lead,
           footnote=ctx.footnote_std("　しきい値は当年度の全庁平均で固定し、前年度も同じしきい値で分類。"))
    x, y, w, h = ctx.content()
    _quadrant_diagram(slide, ctx, x, y, 5.4, h, seg, _seg_rows(ctx, "全庁", "全庁") if not ctx.is_all else None, "全庁")
    rows_lab = [ctx.fy] + ([ctx.prev_fy] if ctx.has_prev else [])
    shares = [seg["share"].tolist()] + ([seg["prev_share"].tolist()] if ctx.has_prev else [])
    p = ch.stacked100(th, ctx.chart("seg_overview"), rows_lab, shares, [lab[k] for k in QUADRANTS],
                      [th.seg_color(k) for k in QUADRANTS], w=3.5, h=1.7, font=7, label_chars=8, show_pct_min=8)
    add_text(slide, th, x + 5.6, y, 3.5, 0.25, "構成比の推移", size=9, bold=True, color=th.hexv("brand_dark"))
    place_image(slide, th, p, x + 5.6, y + 0.27, 3.5, 1.7, align="left")
    desc = [[{"text": lab["high_high"] + "：", "bold": True}, {"text": "意欲・環境ともに良好"}],
            [{"text": lab["high_low"] + "：", "bold": True}, {"text": "意欲は高いが環境上の阻害あり"}],
            [{"text": lab["low_high"] + "：", "bold": True}, {"text": "環境評価は高いが意欲に接続していない"}],
            [{"text": lab["low_low"] + "：", "bold": True}, {"text": "意欲・環境双方に課題"}]]
    bullets(slide, th, x + 5.6, y + 2.1, 3.5, h - 2.1, desc, size=7.5, marker="", spacing=1.1)


@slide_type("segment_by_org")
def s_segment_by_org(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    lab = ctx.cfg.raw["segmentation"]["labels"]
    sg = ctx.res["segments"]
    units = ["全庁"] + [o["name"] for o in ctx.summary["organizations"]]
    shares, ns = [], []
    for u in units:
        d = sg[(sg["scope_type"] == ("全庁" if u == "全庁" else ctx.org_level)) & (sg["scope"] == u)].set_index("segment").reindex(QUADRANTS)
        shares.append(d["share"].tolist())
        ns.append(int(d["n"].iloc[0]) if d["n"].notna().any() else 0)
    all_low = shares[0][3]
    gap = ctx.cfg.analysis["priority_org"]["support_type_gap_pt"]
    hi = [u for u, s in zip(units[1:], shares[1:]) if not np.isnan(s[3]) and s[3] - all_low >= gap]
    lo = [u for u, s in zip(units[1:], shares[1:]) if not np.isnan(s[3]) and all_low - s[3] >= gap]
    lead = f"{ctx.org_level}別の4類型構成比。{lab['low_low']}が全庁（{fmt_pct(all_low)}）比 +{gap}pt 以上: {'・'.join(hi) or 'なし'}。−{gap}pt 以下: {'・'.join(lo) or 'なし'}。"
    _frame(slide, ctx, spec, f"{ctx.org_level}別 職員タイプ構成", lead)
    x, y, w, h = ctx.content()
    p = ch.stacked100(th, ctx.chart("seg_org"), units, shares, [lab[k] for k in QUADRANTS], [th.seg_color(k) for k in QUADRANTS],
                      w=6.3, h=h, font=7, n_labels=ns)
    place_image(slide, th, p, x, y, 6.4, h, align="left")
    items = [[{"text": f"{lab['low_low']}が多い: ", "bold": True, "color": th.hexv("negative")}, {"text": "・".join(hi) or "なし"}],
             [{"text": f"{lab['low_low']}が少ない: ", "bold": True, "color": th.hexv("positive")}, {"text": "・".join(lo) or "なし"}]]
    hh = [u for u, s in zip(units[1:], shares[1:]) if not np.isnan(s[0]) and s[0] - shares[0][0] >= gap]
    items.append([{"text": f"{lab['high_high']}が多い: ", "bold": True, "color": th.hexv("positive")}, {"text": "・".join(hh) or "なし"}])
    add_text(slide, th, x + 6.5, y, 2.6, 0.25, "特徴的な組織（全庁比 ±10pt）", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 6.5, y + 0.3, 2.6, h - 0.3, items, size=7.5, marker="", spacing=1.2)


@slide_type("segment_drivers")
def s_segment_drivers(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    lab = ctx.cfg.raw["segmentation"]["labels"]
    d = ctx.res["segment_drivers"].head(12)
    lead = (f"ドライバー設問の類型別平均。{lab['high_high']}と{lab['low_low']}の差が大きい設問ほど、タイプを分ける要因候補。"
            f"差が最大の設問は「{ctx.qtext(d.iloc[0]['question_id'], 30)}」（{fmt(d.iloc[0]['gap'])}）。" if len(d) else "")
    _frame(slide, ctx, spec, "類型別に差が大きいドライバー設問", lead)
    x, y, w, h = ctx.content()
    cols = [{"id": "q", "title": "設問", "width": 4.4, "kind": "text", "chars": 38},
            {"id": "cat", "title": "カテゴリ", "width": 1.3, "kind": "text", "chars": 11}]
    cols += [{"id": k, "title": lab[k], "width": 0.75, "kind": "num"} for k in QUADRANTS]
    cols.append({"id": "gap", "title": "差", "width": 0.7, "kind": "num", "signed": False})
    rows = [{"cells": {"q": f"{r['question_id']} {r['text']}", "cat": r["category"], **{k: r[k] for k in QUADRANTS}, "gap": r["gap"]}}
            for _, r in d.iterrows()]
    p = ch.render_table(th, ctx.chart("seg_drivers"), cols, rows, row_h=0.3, header_h=0.42, font=7.5)
    place_image(slide, th, p, x, y, w, h, align="left")


# --------------------------------------------------------------------------- T09 Key Driver
def _kd(ctx: Ctx) -> tuple[pd.DataFrame, str]:
    dr = ctx.res["drivers"]
    d = dr[(dr["scope_type"] == ctx.scope_type) & (dr["scope"] == ctx.scope)]
    if len(d):
        return d, ctx.scope_type
    return dr[dr["scope_type"] == "全庁"], "全庁"


def _kd_footnote(ctx: Ctx, src: str) -> str:
    meta = next((m for m in ctx.summary.get("key_driver_meta", []) if m.get("scope") == "全庁"), None)
    s = f"目的変数: {ctx.target_name}（回答者スコア）。説明変数: ドライバー設問（除外: {'・'.join(ctx.cfg.raw['key_driver'].get('exclude', [])) or 'なし'}）。"
    if meta and src == "全庁":
        s += f" 重回帰 n={meta['n']}, 説明変数 {meta['k']}, R²={meta['r2']:.2f}（調整済み {meta['adj_r2']:.2f}）。"
    return s + " 相関・回帰は関連の強さであり因果ではない。"


@slide_type("kd_correlation")
def s_kd_correlation(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    d, src = _kd(ctx)
    top = d.sort_values("r", ascending=False).head(15)
    lead = (f"{ctx.target_name}と各ドライバー設問の相関係数（上位15）。最も関連が強いのは「{ctx.qtext(top.iloc[0]['question_id'], 28)}」（r={fmt(top.iloc[0]['r'])}）。"
            + ("" if src == ctx.scope_type else f"※{ctx.scope}は n<{ctx.cfg.analysis['key_driver']['dept_min_n']} のため全庁の結果。"))
    _frame(slide, ctx, spec, "Key Driver 分析① 相関", lead, footnote=_kd_footnote(ctx, src))
    x, y, w, h = ctx.content()
    labels = [f"{r['question_id']} {ctx.qtext(r['question_id'])}" for _, r in top.iterrows()]
    cols = [th.zone_color(z) for z in top["zone"]]
    p = ch.hbar_scores(th, ctx.chart("kd_corr"), labels, top["r"].tolist(), xlim=(0, 0.7), w=6.4, h=h, colors=cols, font=7,
                       label_chars=22, left=0.42)
    place_image(slide, th, p, x, y, 6.5, h, align="left")
    items = [[{"text": f"{z} {ZONE_LABELS[z]}", "bold": True, "color": th.zone_color(z).lstrip('#')},
              {"text": f"：{int((top['zone'] == z).sum())} 問（上位15内）"}] for z in "ABCD"]
    add_text(slide, th, x + 6.6, y, 2.5, 0.25, "ゾーン（色）", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 6.6, y + 0.3, 2.5, 1.6, items, size=7.5, marker="■ ", spacing=1.1)
    bullets(slide, th, x + 6.6, y + 1.9, 2.5, h - 1.9,
            ["影響度＝相関と標準化回帰係数の順位を合成。上位25%を高影響", "現状スコア＝ドライバー設問平均の中央値で高低を判定",
             "A: 高影響×低スコア＝最優先改善　B: 高影響×高スコア＝重要な強み"], size=7, spacing=1.1)


@slide_type("kd_regression")
def s_kd_regression(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    d, src = _kd(ctx)
    has_beta = d["beta"].notna().any()
    top = (d.sort_values("beta", ascending=False) if has_beta else d.sort_values("r", ascending=False)).head(12)
    if has_beta:
        lead = (f"他の設問を統制した標準化回帰係数（β）の上位12。βが大きい設問ほど、他要因と独立に{ctx.target_name}と結びつく。"
                f"最大は「{ctx.qtext(top.iloc[0]['question_id'], 26)}」（β={fmt(top.iloc[0]['beta'])}）。VIF は多重共線性の目安（10 超で注意）。")
    else:
        lead = f"n が小さいため重回帰は実施せず、相関のみを表示。"
    _frame(slide, ctx, spec, "Key Driver 分析② 重回帰（統制後の関連）", lead, footnote=_kd_footnote(ctx, src))
    x, y, w, h = ctx.content()
    cols = [{"id": "rk", "title": "順位", "width": 0.45, "kind": "num"},
            {"id": "q", "title": "設問", "width": 4.2, "kind": "text", "chars": 40},
            {"id": "cat", "title": "カテゴリ", "width": 1.65, "kind": "text", "chars": 14},
            {"id": "r", "title": "相関 r", "width": 0.6, "kind": "num"},
            {"id": "beta", "title": "β", "width": 0.6, "kind": "num"},
            {"id": "vif", "title": "VIF", "width": 0.5, "kind": "num"},
            {"id": "score", "title": "現状", "width": 0.5, "kind": "num"},
            {"id": "zone", "title": "ゾーン", "width": 0.6, "kind": "text", "chars": 6}]
    rows = []
    for i, (_, r) in enumerate(top.iterrows(), start=1):
        rows.append({"cells": {"rk": i, "q": f"{r['question_id']} {ctx.qtext(r['question_id'])}", "cat": ctx.cfg.q(r["question_id"]).get("category", ""),
                               "r": r["r"], "beta": r["beta"], "vif": fmt(r["vif"], 1), "score": r["score"], "zone": f"{r['zone']}"},
                     "fill": th.zone_color(r["zone"]) + "22" if False else None})
    p = ch.render_table(th, ctx.chart("kd_reg"), cols, rows, row_h=0.3, header_h=0.4, font=7.5)
    place_image(slide, th, p, x, y, w, h, align="left")


@slide_type("kd_impact_score")
def s_kd_impact_score(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    d, src = _kd(ctx)
    a = d[d["zone"] == "A"].sort_values("impact", ascending=False)
    lead = (f"横軸＝現状スコア、縦軸＝影響度（相関・βの合成順位）。右上 B は維持すべき強み、左上 A は改善時の波及が期待される最優先領域。"
            f"A ゾーン {len(a)} 問。")
    _frame(slide, ctx, spec, "Key Driver 分析③ Impact × Score", lead, footnote=_kd_footnote(ctx, src))
    x, y, w, h = ctx.content()
    split = float(d["score_split"].iloc[0])
    thr = float(d[d["high_impact"]]["impact"].min()) if d["high_impact"].any() else float(d["impact"].quantile(0.75))
    cols = [th.zone_color(z) for z in d["zone"]]
    p = ch.bubble(th, ctx.chart("kd_impact"), d["score"].tolist(), d["impact"].tolist(), [1] * len(d), d["question_id"].tolist(),
                  xline=split, yline=thr, xlabel="現状スコア（設問平均）", ylabel="影響度（順位パーセンタイル）", w=6.0, h=h, font=6.5,
                  colors=cols, quadrant_text={"tl": "A 最優先改善", "tr": "B 重要な強み", "bl": "C 二次課題", "br": "D 維持領域"},
                  size_scale=1.0, zone_fill={"tl": th.zone_color("A"), "tr": th.zone_color("B"), "bl": th.zone_color("C"), "br": th.zone_color("D")})
    place_image(slide, th, p, x, y, 6.1, h, align="left")
    items = [[{"text": f"{r['question_id']} ", "bold": True}, {"text": f"{ctx.qtext(r['question_id'])}  r={fmt(r['r'])} / 現状 {fmt(r['score'])}", "size": 7.5}]
             for _, r in a.head(10).iterrows()]
    add_text(slide, th, x + 6.2, y, 2.9, 0.25, "A 最優先改善（高影響 × 低スコア）", size=9, bold=True, color=th.hexv("negative"))
    bullets(slide, th, x + 6.2, y + 0.27, 2.9, h - 0.3, items or ["該当なし"], size=7.5, marker="", spacing=1.05)


# --------------------------------------------------------------------------- T06 / T07 部局
def _org_units(ctx: Ctx) -> list[dict]:
    return ctx.summary["organizations"]


@slide_type("org_overview_table")
def s_org_overview_table(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    units = _org_units(ctx)
    ov = ctx.S[(ctx.S["scope_type"] == ctx.org_level) & (ctx.S["unit_type"] == "指標") & (ctx.S["unit"] == ctx.target)].set_index("scope")
    valid = ov[~ov["suppressed"]]
    hi, lo = valid["mean"].idxmax(), valid["mean"].idxmin()
    lead = (f"{ctx.org_level}別の{ctx.target_name}。最高は{hi}（{fmt(valid.loc[hi, 'mean'])}）、最低は{lo}（{fmt(valid.loc[lo, 'mean'])}）。"
            f"全庁差 ±{ctx.cfg.analysis['judgement']['org_gap_abs_at_least']} 以上の組織を特徴とみなす。青／桃は有意差（Welch t）。")
    _frame(slide, ctx, spec, f"{ctx.org_level}別 {ctx.target_name} 一覧", lead)
    x, y, w, h = ctx.content()
    axes = ctx.summary["axes"]
    cols = [{"id": "org", "title": ctx.org_level, "width": 1.8, "kind": "text", "chars": 14},
            {"id": "n", "title": "n", "width": 0.5, "kind": "num"},
            {"id": "mean", "title": ctx.target_name, "width": 0.9, "kind": "num"},
            {"id": "gap", "title": "全庁差", "width": 0.8, "kind": "sig"},
            {"id": "yoy", "title": "前年差", "width": 0.8, "kind": "sig"}]
    cols += [{"id": f"ax{i}", "title": a, "width": 0.8, "kind": "num"} for i, a in enumerate(axes)]
    cols += [{"id": "bar", "title": "肯定/中立/否定", "width": w - sum(c["width"] for c in cols), "kind": "bar3"}]
    rows = []
    allr = ctx.row("指標", ctx.target, "全庁", "全庁")
    rows.append({"style": "total", "cells": {"org": "全庁", "n": int(allr["n"]), "mean": allr["mean"], "gap": ("", 0),
                                              "yoy": _sig_tuple(allr["diff_yoy"], allr["sig_yoy"]),
                                              **{f"ax{i}": ctx.val("軸", a, "mean", "全庁", "全庁") for i, a in enumerate(axes)},
                                              "bar": (allr["pos"], allr["neu"], allr["neg"])}})
    for u in units:
        r = ov.loc[u["name"]] if u["name"] in ov.index else None
        if r is None or r["suppressed"]:
            rows.append({"cells": {"org": u["name"], "n": u["n"], "mean": "－", "gap": ("－", 0), "yoy": ("－", 0), "bar": None,
                                   **{f"ax{i}": "－" for i in range(len(axes))}}})
            continue
        rows.append({"cells": {"org": u["name"], "n": int(r["n"]), "mean": r["mean"], "gap": _sig_tuple(r["diff_all"], r["sig_all"]),
                               "yoy": _sig_tuple(r["diff_yoy"], r["sig_yoy"]),
                               **{f"ax{i}": ctx.val("軸", a, "mean", ctx.org_level, u["name"]) for i, a in enumerate(axes)},
                               "bar": (r["pos"], r["neu"], r["neg"])}})
    rh = min(0.26, (h - 0.4) / len(rows))
    p = ch.render_table(th, ctx.chart("org_table"), cols, rows, row_h=rh, header_h=0.38, font=7 if rh >= 0.22 else 6.3)
    place_image(slide, th, p, x, y, w, h, align="left")


@slide_type("org_heatmap")
def s_org_heatmap(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    value = spec.get("value", "diff_all")
    is_yoy = value == "diff_yoy"
    if is_yoy and not ctx.has_prev:
        _frame(slide, ctx, spec, f"{ctx.org_level} × カテゴリ（前年差）", "前年度データが無いため作成していない。")
        return
    units = _org_units(ctx)
    cats = ctx.summary["categories"]
    S = ctx.S
    d = S[(S["scope_type"] == ctx.org_level) & (S["unit_type"] == "カテゴリ")]
    mat = d.pivot(index="scope", columns="unit", values=value).reindex(index=[u["name"] for u in units], columns=cats)
    tot = S[(S["scope_type"] == ctx.org_level) & (S["unit_type"] == "指標") & (S["unit"] == ctx.target)].set_index("scope")["mean"].reindex(mat.index)
    thr = ctx.cfg.analysis["judgement"]["org_gap_abs_at_least" if not is_yoy else "improved_at_least"]
    neg = int((mat.values <= -thr).sum())
    pos = int((mat.values >= thr).sum())
    title = f"{ctx.org_level} × カテゴリ ヒートマップ（{'前年差' if is_yoy else '全庁差'}）"
    lead = (f"セルは{ctx.org_level}のカテゴリ平均 − {'前年' if is_yoy else '全庁'}平均。青＝高い／改善、桃＝低い／悪化。"
            f"±{thr} 以上のセル: プラス {pos}、マイナス {neg}。左列は{ctx.target_name}の絶対値。")
    _frame(slide, ctx, spec, title, lead, footnote=ctx.footnote_std("　前年差は比較可能なカテゴリのみ（新設設問を含むカテゴリは参考比較）。"))
    x, y, w, h = ctx.content()
    p = ch.heatmap(th, ctx.chart(f"org_heat_{value}"), list(mat.index), [short_cat(c) for c in cats], mat.values.astype(float),
                   vlim=0.3, w=w, h=h, font=6.8, row_chars=11, col_chars=6, n_col=[u["n"] for u in units],
                   cbar_label="前年差" if is_yoy else "全庁差", first_col_values=tot.values, first_col_title="総合")
    place_image(slide, th, p, x, y, w, h, align="left")


@slide_type("org_positioning")
def s_org_positioning(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    pc = ctx.cfg.raw["positioning"]
    xs_key, ys_key = pc["x"], pc["y"]

    def _ut(k):
        return "指標" if k in ctx.cfg.indicators else "軸"

    units = [u for u in _org_units(ctx) if not u["suppressed"]]
    xs = [ctx.val(_ut(xs_key), xs_key, "mean", ctx.org_level, u["name"]) for u in units]
    ys = [ctx.val(_ut(ys_key), ys_key, "mean", ctx.org_level, u["name"]) for u in units]
    xl = ctx.val(_ut(xs_key), xs_key, "mean", "全庁", "全庁")
    yl = ctx.val(_ut(ys_key), ys_key, "mean", "全庁", "全庁")
    xn, yn = ctx.cfg.indicator_name(xs_key) if xs_key in ctx.cfg.indicators else xs_key, ctx.cfg.indicator_name(ys_key) if ys_key in ctx.cfg.indicators else ys_key
    tr = [u["name"] for u, a, b in zip(units, xs, ys) if a >= xl and b >= yl]
    bl = [u["name"] for u, a, b in zip(units, xs, ys) if a < xl and b < yl]
    lead = f"{ctx.org_level}を{xn}（横）×{yn}（縦）に配置。バブルの大きさ＝回答者数、点線＝全庁平均。右上: {'・'.join(tr) or 'なし'}。左下: {'・'.join(bl) or 'なし'}。"
    _frame(slide, ctx, spec, f"{ctx.org_level}ポジショニング", lead)
    x, y, w, h = ctx.content()
    cols = [th.c("positive") if (a >= xl and b >= yl) else th.c("negative") if (a < xl and b < yl) else th.c("brand") for a, b in zip(xs, ys)]
    p = ch.bubble(th, ctx.chart("org_pos"), xs, ys, [u["n"] for u in units], [u["name"] for u in units], xline=xl, yline=yl,
                  xlabel=f"{xn}（平均）", ylabel=f"{yn}（平均）", w=6.4, h=h, font=7, colors=cols,
                  quadrant_text={"tr": "両方高い", "bl": "両方低い", "tl": f"{yn}のみ高い", "br": f"{xn}のみ高い"})
    place_image(slide, th, p, x, y, 6.5, h, align="left")
    items = [[{"text": "右上（両方高い）: ", "bold": True, "color": th.hexv("positive")}, {"text": "・".join(tr) or "なし"}],
             [{"text": "左下（両方低い）: ", "bold": True, "color": th.hexv("negative")}, {"text": "・".join(bl) or "なし"}],
             [{"text": "全庁平均: ", "bold": True}, {"text": f"{xn} {fmt(xl)} / {yn} {fmt(yl)}"}]]
    add_text(slide, th, x + 6.6, y, 2.5, 0.25, "位置づけ", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 6.6, y + 0.3, 2.5, h - 0.3, items, size=7.5, marker="", spacing=1.2)


@slide_type("priority_orgs")
def s_priority_orgs(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    pr = ctx.res["priority_orgs"]
    pts = ctx.cfg.analysis["priority_org"]["points"]
    bands = ctx.cfg.analysis["priority_org"]["bands"]
    focus = pr[pr["band"] == "重点フォロー候補"]["unit"].tolist()
    good = pr[pr["good_practice"]]["unit"].tolist()
    lead = (f"6条件の加点で重点フォロー候補を抽出（{bands['priority']}点以上＝重点フォロー候補、{bands['watch']}〜{bands['priority'] - 1}点＝要注視）。"
            f"候補: {'・'.join(focus) or 'なし'}。好事例候補（高水準×改善）: {'・'.join(good) or 'なし'}。最終確定は CAQNAL レビュー。")
    _frame(slide, ctx, spec, f"重点{ctx.org_level}・好事例候補", lead,
           footnote="加点: 総合が全庁下位20% +2 / 前年差悪化 +2 / 全庁差≤−0.15 +2 / Key Driver 上位項目が低評価 +2 / 自由記述で同テーマ多い +1 / 要支援型が全庁比+10pt +1")
    x, y, w, h = ctx.content()
    labels = {"overall_bottom_pct": "下位20%", "worsened": "悪化", "org_gap_negative": "全庁差−", "driver_low": "KD低評価",
              "comment_theme": "自由記述", "support_type_gap": "要支援型+"}
    cols = [{"id": "org", "title": ctx.org_level, "width": 1.5, "kind": "text", "chars": 12},
            {"id": "n", "title": "n", "width": 0.4, "kind": "num"},
            {"id": "mean", "title": "総合", "width": 0.6, "kind": "num"},
            {"id": "gap", "title": "全庁差", "width": 0.6, "kind": "num", "signed": True},
            {"id": "yoy", "title": "前年差", "width": 0.6, "kind": "num", "signed": True}]
    cols += [{"id": k, "title": f"{labels[k]}(+{pts[k]})", "width": 0.55, "kind": "text", "chars": 6, "align": "center"} for k in pts]
    cols += [{"id": "total", "title": "合計", "width": 0.45, "kind": "num"},
             {"id": "band", "title": "区分", "width": 1.0, "kind": "text", "chars": 9},
             {"id": "good", "title": "好事例", "width": w - 1.5 - 0.4 - 0.6 * 3 - 0.55 * len(pts) - 0.45 - 1.0, "kind": "text", "chars": 4, "align": "center"}]
    rows = []
    for _, r in pr.iterrows():
        fill = th.c("negative_soft") if r["band"] == "重点フォロー候補" else th.c("brand_pale") if r["band"] == "要注視" else None
        cells = {"org": r["unit"], "n": int(r["n"]), "mean": r["mean"], "gap": r["diff_all"], "yoy": r["diff_yoy"],
                 "total": int(r["total"]), "band": r["band"], "good": "○" if r["good_practice"] else ""}
        for k in pts:
            cells[k] = f"+{int(r[k])}" if r[k] else ""
        rows.append({"cells": cells, "fill": fill})
    rh = min(0.25, (h - 0.4) / max(len(rows), 1))
    p = ch.render_table(th, ctx.chart("priority_orgs"), cols, rows, row_h=rh, header_h=0.42, font=6.8, zebra=False)
    place_image(slide, th, p, x, y, w, h, align="left")


# --------------------------------------------------------------------------- T10 属性
@slide_type("attribute_heatmap")
def s_attribute_heatmap(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    attrs = spec.get("attributes") or list(ctx.cfg.attributes)[:2]
    cats = ctx.summary["categories"]
    S = ctx.S
    thr = ctx.cfg.analysis["judgement"]["org_gap_abs_at_least"]
    _frame(slide, ctx, spec, f"属性別結果（{'・'.join(attrs)}）",
           f"属性区分ごとのカテゴリ平均 − 全庁平均。青＝高い、桃＝低い（±{thr} 以上を特徴とみなす）。左列は{ctx.target_name}の絶対値。")
    x, y, w, h = ctx.content()
    n = len(attrs)
    each = (h - 0.1 * (n - 1)) / n
    for i, key in enumerate(attrs):
        spec_a = ctx.cfg.attributes[key]
        order = list(spec_a.get("order", []))
        short = spec_a.get("short", {})
        d = S[(S["scope_type"] == key) & (S["unit_type"] == "カテゴリ")]
        vals = [v for v in order if v in set(d["scope"])] + sorted(set(d["scope"]) - set(order))
        mat = d.pivot(index="scope", columns="unit", values="diff_all").reindex(index=vals, columns=cats)
        tot = S[(S["scope_type"] == key) & (S["unit_type"] == "指標") & (S["unit"] == ctx.target)].set_index("scope").reindex(vals)
        cy = y + i * (each + 0.1)
        p = ch.heatmap(th, ctx.chart(f"attr_heat_{i}"), [short.get(v, v) for v in vals], [short_cat(c) for c in cats],
                       mat.values.astype(float), vlim=0.3, w=w, h=max(each, 1.2), font=6.8, row_chars=11, col_chars=6,
                       n_col=tot["n"].fillna(0).astype(int).tolist(), cbar_label="全庁差", first_col_values=tot["mean"].values,
                       first_col_title="総合", title=key)
        place_image(slide, th, p, x, cy, w, each, align="left")


# --------------------------------------------------------------------------- T11 自由記述
@slide_type("comments_overview")
def s_comments_overview(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    ct = ctx.res["comment_themes"]
    if not len(ct):
        _frame(slide, ctx, spec, "自由記述の概況", "自由記述設問が無いため作成していない。")
        return
    d = ct[(ct["scope_type"] == ctx.scope_type) & (ct["scope"] == ctx.scope)]
    total = int(d["total_comments"].iloc[0]) if len(d) else 0
    themes = d[~d["theme"].str.startswith("トーン:")].sort_values("count", ascending=False)
    tones = d[d["theme"].str.startswith("トーン:")]
    cm = ctx.res["comments"]
    avg_len = float(cm["raw_len"].mean()) if len(cm) else float("nan")
    top = themes.iloc[0] if len(themes) else None
    lead = (f"自由記述 {total} 件（回答者の {100 * total / max(ctx.n_scope(), 1):.0f}%、平均 {fmt(avg_len, 0)} 字）。"
            + (f"最多テーマは「{top['theme']}」{int(top['count'])} 件（{fmt_pct(top['share'])}）。" if top is not None and not np.isnan(top["count"]) else "")
            + "キーワード辞書による規則分類（複数テーマ可）。")
    _frame(slide, ctx, spec, "自由記述の概況（テーマ構成）", lead,
           footnote=f"件数 {ctx.cfg.analysis['anonymity']['min_comments']} 未満は非表示。個人特定情報はマスキング済み。原文はレポートに掲載しない。トーンは肯定・否定語の出現数比較による参考値。")
    x, y, w, h = ctx.content()
    tl = themes[themes["theme"] != "その他"]
    tl = pd.concat([tl, themes[themes["theme"] == "その他"]])
    p = ch.bar_counts(th, ctx.chart("comment_themes"), tl["theme"].tolist(), tl["count"].tolist(), w=5.6, h=h, font=8, label_chars=14, pct=True, total=total)
    place_image(slide, th, p, x, y, 5.7, h, align="left")
    tone_items = [[{"text": f"{r['theme'].split(':')[1]}: ", "bold": True}, {"text": f"{int(r['count']) if not np.isnan(r['count']) else '－'} 件（{fmt_pct(r['share'])}）"}] for _, r in tones.iterrows()]
    add_text(slide, th, x + 5.9, y, 3.2, 0.25, "トーン（参考）", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 5.9, y + 0.27, 3.2, 1.1, tone_items or ["－"], size=7.5, marker="", spacing=1.1)
    # 定量接続
    inv: dict[str, list[str]] = {}
    cats_here = set(ctx.summary["categories"])
    for cat, t in (ctx.cfg.action_library.get("category_theme_map") or {}).items():
        if t and cat in cats_here:            # 他自治体向けの既定マップに含まれるカテゴリは無視
            inv.setdefault(t, []).append(cat)
    items = []
    for _, r in themes.head(4).iterrows():
        if r["theme"] == "その他" or np.isnan(r["count"]):
            continue
        cats = inv.get(r["theme"], [])
        if cats:
            means = [f"{c} {fmt(ctx.val('カテゴリ', c))}" for c in cats[:2]]
            items.append([{"text": f"{r['theme']}（{int(r['count'])}件）→ ", "bold": True}, {"text": "関連カテゴリ: " + "／".join(means)}])
    add_text(slide, th, x + 5.9, y + 1.5, 3.2, 0.25, "定量結果との接続", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 5.9, y + 1.77, 3.2, h - 1.8, items or ["該当なし"], size=7.5, marker="", spacing=1.1)


@slide_type("comments_by_org")
def s_comments_by_org(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    ct = ctx.res["comment_themes"]
    if not len(ct):
        _frame(slide, ctx, spec, f"{ctx.org_level}別 自由記述テーマ", "自由記述設問が無いため作成していない。")
        return
    minc = ctx.cfg.analysis["anonymity"]["min_comments"]
    _frame(slide, ctx, spec, f"{ctx.org_level}別 自由記述テーマ",
           f"{ctx.org_level}ごとの件数と上位テーマ。件数 {minc} 未満は非表示。少数コメントからの一般化は避け、定量結果の補足として読む。")
    x, y, w, h = ctx.content()
    cols = [{"id": "org", "title": ctx.org_level, "width": 1.8, "kind": "text", "chars": 14},
            {"id": "n", "title": "件数", "width": 0.6, "kind": "num"},
            {"id": "t1", "title": "テーマ 1位", "width": 2.2, "kind": "text", "chars": 18},
            {"id": "t2", "title": "テーマ 2位", "width": 2.2, "kind": "text", "chars": 18},
            {"id": "t3", "title": "テーマ 3位", "width": w - 6.8, "kind": "text", "chars": 18}]
    rows = []
    for stype, scope in [("全庁", "全庁")] + [(ctx.org_level, o["name"]) for o in _org_units(ctx)]:
        d = ct[(ct["scope_type"] == stype) & (ct["scope"] == scope) & (~ct["theme"].str.startswith("トーン:")) & (ct["theme"] != "その他")]
        total = int(d["total_comments"].iloc[0]) if len(d) else 0
        if not len(d) or bool(d["suppressed"].iloc[0]):
            rows.append({"cells": {"org": scope, "n": total, "t1": "－（非表示）", "t2": "", "t3": ""}})
            continue
        top = d.sort_values("count", ascending=False).head(3)
        cells = {"org": scope, "n": total}
        for i, (_, r) in enumerate(top.iterrows(), start=1):
            cells[f"t{i}"] = f"{r['theme']}（{int(r['count'])}件・{fmt_pct(r['share'])}）" if r["count"] > 0 else ""
        rows.append({"cells": cells, "style": "total" if stype == "全庁" else "item"})
    rh = min(0.25, (h - 0.4) / len(rows))
    p = ch.render_table(th, ctx.chart("comments_org"), cols, rows, row_h=rh, header_h=0.36, font=6.8)
    place_image(slide, th, p, x, y, w, h, align="left")


# --------------------------------------------------------------------------- T14 総括 / 重点課題
@slide_type("synthesis")
def s_synthesis(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    ins = ctx.ins() or {}
    sen = ins.get("sentences", {})
    _frame(slide, ctx, spec, "分析の総括", sen.get("fact", ""),
           footnote=ctx.footnote_std("　示唆は数値から規則的に生成した初期仮説。重点課題・施策の最終確定は CAQNAL レビューで行う。"))
    x, y, w, h = ctx.content()
    colw = (w - 0.2) / 3
    jd = ctx.res["judgement"]

    def block(bx, by, bw, bh, title, items, color):
        add_rect(slide, th, bx, by, bw, 0.26, color)
        add_text(slide, th, bx, by, bw, 0.26, title, size=9.5, bold=True, color="FFFFFF", anchor="middle")
        add_rect(slide, th, bx, by + 0.26, bw, bh - 0.26, th.hexv("bg"), th.hexv("line"))
        bullets(slide, th, bx + 0.05, by + 0.3, bw - 0.1, bh - 0.35, items, size=7.5, marker="", spacing=1.05)

    st = jd[jd["strength"]].sort_values("mean", ascending=False).head(5)
    items = [[{"text": f"{r['question_id']} ", "bold": True}, {"text": f"{r['text']}"},
              {"text": f"　平均 {fmt(r['mean'])}・肯定率 {fmt_pct(r['pos'])}" + (f"・前年差 {fmt(r['diff_yoy'], signed=True)}" if not np.isnan(r["diff_yoy"]) else ""), "size": 6.8, "color": th.hexv("ink_soft")}]
             for _, r in st.iterrows()]
    block(x, y, colw, 2.6, "強み（高評価 × 悪化なし）", items or ["該当なし"], th.hexv("positive"))
    pi = jd[jd["priority_issue"]].head(5)
    items = [[{"text": f"{r['question_id']} ", "bold": True}, {"text": f"{r['text']}"},
              {"text": f"　{r['signal_labels']}（{int(r['signals'])}観点）・平均 {fmt(r['mean'])}", "size": 6.8, "color": th.hexv("ink_soft")}]
             for _, r in pi.iterrows()]
    block(x + colw + 0.1, y, colw, 2.6, "重点課題候補（3観点以上）", items or ["該当なし"], th.hexv("negative"))
    kd = [k for k in ins.get("key_drivers", []) if k["zone"] == "A"][:5] or ins.get("key_drivers", [])[:5]
    items = [[{"text": f"{k['question_id']} ", "bold": True}, {"text": f"{k['text']}"},
              {"text": f"　r={fmt(k['r'])}" + (f"・β={fmt(k['beta'])}" if k.get("beta") is not None else "") + f"・現状 {fmt(k['score'])}・{k['zone']}", "size": 6.8, "color": th.hexv("ink_soft")}]
             for k in kd]
    block(x + 2 * (colw + 0.1), y, colw, 2.6, "主要因候補（Key Driver A ゾーン）", items or ["該当なし"], th.hexv("brand_dark"))
    # 下段: 重点組織・層
    y2 = y + 2.7
    pr = ctx.res["priority_orgs"]
    focus = pr[pr["band"] == "重点フォロー候補"]["unit"].tolist() if len(pr) else []
    watch = pr[pr["band"] == "要注視"]["unit"].tolist() if len(pr) else []
    # 属性で全庁差が大きいマイナス
    S = ctx.S
    at = S[(S["scope_type"].isin(list(ctx.cfg.attributes))) & (S["unit_type"] == "指標") & (S["unit"] == ctx.target) & (~S["suppressed"])]
    at = at.sort_values("diff_all").head(3)
    layers = [f"{r['scope_type']}＝{r['scope']}（全庁差 {fmt(r['diff_all'], signed=True)}）" for _, r in at.iterrows() if r["diff_all"] < 0]
    items = [[{"text": f"重点フォロー候補{ctx.org_level}: ", "bold": True}, {"text": "・".join(focus) or "なし"}],
             [{"text": f"要注視{ctx.org_level}: ", "bold": True}, {"text": "・".join(watch) or "なし"}],
             [{"text": f"{ctx.target_name}が相対的に低い層: ", "bold": True}, {"text": "／".join(layers) or "なし"}]]
    block(x, y2, w * 0.48, h - 2.7, "重点組織・層", items, th.hexv("ink_soft"))
    items = [sen.get("feature", ""), sen.get("cause", ""), sen.get("hint", "")]
    block(x + w * 0.48 + 0.1, y2, w * 0.52 - 0.1, h - 2.7, "所見（規則生成）", [i for i in items if i], th.hexv("brand"))


@slide_type("priority_issues")
def s_priority_issues(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    jd = ctx.res["judgement"]
    d = jd[jd["signals"] >= 2].head(14)
    n_pri = int(jd["priority_issue"].sum())
    lead = (f"6観点（低スコア・悪化・組織偏在・Key Driver・自由記述・ばらつき）のシグナルを設問ごとに集計。"
            f"3観点以上の重点課題候補は {n_pri} 問。●＝該当。")
    _frame(slide, ctx, spec, "重点課題候補（6観点の統合）", lead,
           footnote="低スコア: 全庁下位20%または否定率20%以上／悪化: 前年差−0.10以下／組織偏在: いずれかの部局で全庁差−0.15以下／Key Driver: 影響度上位25%／自由記述: 関連テーマ5件以上／ばらつき: SD上位20%")
    x, y, w, h = ctx.content()
    cols = [{"id": "q", "title": "設問", "width": 4.2, "kind": "text", "chars": 36},
            {"id": "cat", "title": "カテゴリ", "width": 1.2, "kind": "text", "chars": 10},
            {"id": "mean", "title": "平均", "width": 0.5, "kind": "num"},
            {"id": "neg", "title": "否定率", "width": 0.55, "kind": "pct"},
            {"id": "yoy", "title": "前年差", "width": 0.6, "kind": "num", "signed": True}]
    cols += [{"id": s, "title": SIGNAL_LABELS[s], "width": 0.5, "kind": "text", "chars": 5} for s in SIGNALS]
    cols += [{"id": "sig", "title": "計", "width": w - sum(c["width"] for c in cols), "kind": "num"}]
    rows = []
    for _, r in d.iterrows():
        cells = {"q": f"{r['question_id']} {r['text']}", "cat": r["category"], "mean": r["mean"], "neg": r["neg"], "yoy": r["diff_yoy"], "sig": int(r["signals"])}
        for s in SIGNALS:
            cells[s] = "●" if r[s] else ""
        rows.append({"cells": cells, "fill": th.c("negative_soft") if r["priority_issue"] else None})
    rh = min(0.28, (h - 0.4) / max(len(rows), 1))
    p = ch.render_table(th, ctx.chart("priority_issues"), cols, rows, row_h=rh, header_h=0.4, font=7, zebra=False)
    place_image(slide, th, p, x, y, w, h, align="left")


# --------------------------------------------------------------------------- T15 施策
def _issue_themes(ctx: Ctx, n: int = 3) -> list[tuple[str, list[str]]]:
    """課題設問 → テーマ（上位 n）と根拠設問。"""
    ins = ctx.ins() or {}
    cmap = ctx.cfg.action_library.get("category_theme_map") or {}
    if ctx.is_all:
        jd = ctx.res["judgement"]
        src = jd[jd["signals"] >= 1].sort_values(["signals", "mean"], ascending=[False, True])
        pairs = [(cmap.get(r["category"]), f"{r['question_id']} {r['text']}") for _, r in src.iterrows()]
    else:
        pairs = [(cmap.get(i["category"]), f"{i['question_id']} {i['text']}") for i in ins.get("top_issues", [])]
        rk = _ranking_rows(ctx, "全庁比マイナス", 5)
        pairs += [(cmap.get(ctx.cfg.q(q).get("category")), f"{q} {ctx.qtext(q)}") for q in rk["question_id"]]
    out: dict[str, list[str]] = {}
    for t, q in pairs:
        if t:
            out.setdefault(t, [])
            if q not in out[t]:
                out[t].append(q)
    return list(out.items())[:n]


@slide_type("actions")
def s_actions(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    lib = ctx.cfg.action_library.get("themes", {})
    themes = _issue_themes(ctx, 4)
    _frame(slide, ctx, spec, "施策の方向性（課題テーマ → 現場確認 → 施策候補）",
           "課題テーマごとに、現場で確認すべき要因と施策の選択肢を提示する。施策は確定ではなく検討候補であり、管理職WSと人事検討で具体化する。",
           footnote="施策優先度は Impact（関連度・対象人数・深刻度）× Ease（コスト・権限・期間）で別途評価。短期=3〜6か月／中期=6〜18か月／長期=18か月以上。")
    x, y, w, h = ctx.content()
    cols = [{"id": "theme", "title": "課題テーマ", "width": 1.3, "kind": "text", "chars": 8, "lines": 2},
            {"id": "ev", "title": "根拠となる課題設問", "width": 2.6, "kind": "text", "chars": 20, "lines": 3},
            {"id": "check", "title": "現場で確認すべき要因", "width": 1.6, "kind": "text", "chars": 12, "lines": 3},
            {"id": "act", "title": "施策候補", "width": 1.7, "kind": "text", "chars": 13, "lines": 3},
            {"id": "lv", "title": "全庁／部局／課の役割", "width": w - 7.2, "kind": "text", "chars": 15, "lines": 3}]
    rows = []
    for t, qs in themes:
        L = lib.get(t, {})
        lv = L.get("level", {})
        rows.append({"cells": {"theme": t, "ev": "\n".join(qs[:2]), "check": "・".join(L.get("check", [])),
                               "act": "・".join(L.get("actions", [])),
                               "lv": f"全庁: {lv.get('全庁', '')}／部局: {lv.get('部局', '')}／課: {lv.get('課', '')}"}})
    p = ch.render_table(th, ctx.chart("actions"), cols, rows or [{"cells": {"theme": "該当なし"}}], row_h=min(0.95, (h - 0.5) / max(len(rows), 1)),
                        header_h=0.4, font=7.2, zebra=True)
    place_image(slide, th, p, x, y, w, h, align="left")


# --------------------------------------------------------------------------- APPENDIX
@slide_type("question_list")
def s_question_list(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    qs = ctx.cfg.qs(("outcome", "driver"))
    part, parts = spec.get("part", 1), spec.get("parts", 2)
    per = int(np.ceil(len(qs) / parts))
    sub = qs[(part - 1) * per: part * per]
    _frame(slide, ctx, spec, f"設問一覧と全庁結果（{part}/{parts}）", "尺度設問の一覧。逆転項目は分析スコアを反転して集計。※は回答者限定の設問。", footnote=ctx.footnote_std())
    x, y, w, h = ctx.content()
    cols = [{"id": "id", "title": "ID", "width": 0.45, "kind": "text", "chars": 4},
            {"id": "axis", "title": "軸", "width": 0.8, "kind": "text", "chars": 6},
            {"id": "cat", "title": "カテゴリ", "width": 1.4, "kind": "text", "chars": 12},
            {"id": "q", "title": "設問", "width": 4.0, "kind": "text", "chars": 38},
            {"id": "mean", "title": "平均", "width": 0.55, "kind": "num"},
            {"id": "pos", "title": "肯定率", "width": 0.6, "kind": "pct"},
            {"id": "neg", "title": "否定率", "width": 0.6, "kind": "pct"},
            {"id": "yoy", "title": "前年差", "width": w - 8.4, "kind": "sig"}]
    rows = []
    for q in sub:
        r = ctx.row("設問", q["id"], "全庁", "全庁")
        flag = ("※" if q.get("condition") else "") + ("（逆転）" if q.get("reverse") else "")
        rows.append({"cells": {"id": q["id"], "axis": q.get("axis", ""), "cat": q.get("category", ""), "q": q["text"] + flag,
                               "mean": r["mean"] if r is not None else None, "pos": r["pos"] if r is not None else None,
                               "neg": r["neg"] if r is not None else None,
                               "yoy": _sig_tuple(r["diff_yoy"], r["sig_yoy"]) if r is not None else ("－", 0)}})
    rh = (h - 0.3) / max(len(rows), 1)
    p = ch.render_table(th, ctx.chart(f"qlist_{part}"), cols, rows, row_h=min(rh, 0.24), header_h=0.3, font=6.3)
    place_image(slide, th, p, x, y, w, h, align="left")


@slide_type("method_notes")
def s_method_notes(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    a = ctx.cfg.analysis
    _frame(slide, ctx, spec, "分析手法・注記", "本レポートの計算方法と前提。設定値は settings_snapshot.yaml として納品データに含める。", footnote="")
    x, y, w, h = ctx.content()
    rev = [q["id"] for q in ctx.cfg.questions if q.get("reverse")]
    cond = [f"{q['id']}（{q['condition']['attribute']}={q['condition']['equals']} のみ）" for q in ctx.cfg.questions if q.get("condition")]
    left = [
        [{"text": "スコア化　", "bold": True}, {"text": "選択肢を 5〜1 に数値化。逆転項目: " + ("・".join(rev) or "なし")}],
        [{"text": "条件付き設問　", "bold": True}, {"text": "／".join(cond) or "なし"}],
        [{"text": "複合単位　", "bold": True}, {"text": f"カテゴリ・軸・指標は構成設問の回答者平均（指標は構成設問の {int(a['score']['indicator_min_answer_ratio'] * 100)}% 以上回答が必要）"}],
        [{"text": "有意差　", "bold": True}, {"text": f"平均差は Welch の t 検定、α={a['significance']['alpha']}。多重比較の補正は行っていない"}],
        [{"text": "Key Driver　", "bold": True}, {"text": f"Pearson 相関（ペア n≥{a['key_driver']['min_n']}）と標準化重回帰（OLS、n≥説明変数×{a['key_driver']['ols_min_ratio']}）。VIF で多重共線性を確認"}],
        [{"text": "4類型　", "bold": True}, {"text": f"しきい値={a['segmentation']['threshold']}（当年度全庁）。前年度も同じしきい値で分類"}],
    ]
    bullets(slide, th, x, y, 4.5, h, left, size=7.5, spacing=1.15)
    j = a["judgement"]
    right = [
        [{"text": "高評価／低評価　", "bold": True}, {"text": f"全庁上位{j['high']['top_pct']}%または肯定率{j['high']['positive_pct_at_least']}%以上／下位{j['low']['bottom_pct']}%または否定率{j['low']['negative_pct_at_least']}%以上"}],
        [{"text": "改善／悪化　", "bold": True}, {"text": f"前年差 {j['improved_at_least']:+.2f} 以上／{j['worsened_at_most']:+.2f} 以下"}],
        [{"text": "全庁比特徴　", "bold": True}, {"text": f"全庁差 ±{j['org_gap_abs_at_least']} 以上"}],
        [{"text": "高影響／ばらつき大　", "bold": True}, {"text": f"Key Driver 上位{j['high_impact_top_pct']}%／SD 上位{j['high_sd_top_pct']}%"}],
        [{"text": "重点課題候補　", "bold": True}, {"text": f"6観点のうち {j['priority_issue_min_signals']} 以上"}],
        [{"text": "匿名性　", "bold": True}, {"text": f"n<{a['anonymity']['min_n']} の組織・属性、{a['anonymity']['min_comments']} 件未満の自由記述は非表示。二次秘匿は CAQNAL と協議"}],
        [{"text": "経年比較　", "bold": True}, {"text": "設問文の一致で対応付け。文言変更は参考比較、新設は比較不可（yoy_map.yaml）"}],
    ]
    bullets(slide, th, x + 4.7, y, 4.4, h, right, size=7.5, spacing=1.15)


# --------------------------------------------------------------------------- 部局別 (T12)
@slide_type("dept_summary")
def s_dept_summary(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    ins = ctx.ins() or {}
    sen = ins.get("sentences", {})
    ov = ctx.row("指標", ctx.target)
    _frame(slide, ctx, spec, f"{ctx.scope}｜サマリー", sen.get("fact", ""))
    x, y, w, h = ctx.content()
    cw = (w - 0.3) / 4
    kpi_card(slide, th, x, y, cw, 1.0, ctx.target_name, fmt(ov["mean"]),
             f"全庁 {fmt(ctx.val('指標', ctx.target, 'mean', '全庁', '全庁'))}（差 {fmt(ov['diff_all'], signed=True)}{'・有意' if ov['sig_all'] else ''}）")
    kpi_card(slide, th, x + cw + 0.1, y, cw, 1.0, "前年差", fmt(ov["diff_yoy"], signed=True) + _ref_mark(ov["yoy_status"]) if not np.isnan(ov["diff_yoy"]) else "－",
             f"前年 {fmt(ov['prev_mean'])}（n={int(ov['prev_n'])}）" if not np.isnan(ov["prev_mean"]) else "前年データなし／非表示",
             accent=th.hexv("positive") if ov["diff_yoy"] >= 0 else th.hexv("negative"))
    for i, ax in enumerate(ctx.summary["axes"][:2]):
        r = ctx.row("軸", ax)
        kpi_card(slide, th, x + (2 + i) * (cw + 0.1), y, cw, 1.0, ax, fmt(r["mean"]),
                 f"全庁差 {fmt(r['diff_all'], signed=True)} / 前年差 {fmt(r['diff_yoy'], signed=True)}", accent=th.hexv("brand"))
    cats = ctx.summary["categories"]
    rows = ctx.rows("カテゴリ").set_index("unit").reindex(cats)
    ref_all = [ctx.val("カテゴリ", c, "mean", "全庁", "全庁") for c in cats]
    p = ch.hbar_scores(th, ctx.chart("dept_cat"), cats, rows["mean"].tolist(), prev=ref_all, diffs=rows["diff_all"].tolist(),
                       diff_title="全庁差", sig=rows["sig_all"].tolist(), w=5.6, h=h - 1.15, font=7.5, label_chars=14,
                       colors=[th.c("positive") if (not np.isnan(d) and d >= 0) else th.c("negative") for d in rows["diff_all"]])
    add_text(slide, th, x, y + 1.1, 5.6, 0.25, "カテゴリ別平均（◇＝全庁）", size=9, bold=True, color=th.hexv("brand_dark"))
    place_image(slide, th, p, x, y + 1.35, 5.7, h - 1.4, align="left")
    items = [s for s in (sen.get("feature"), sen.get("cause"), sen.get("hint")) if s]
    add_text(slide, th, x + 5.9, y + 1.1, 3.2, 0.25, "所見（規則生成・初期仮説）", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 5.9, y + 1.37, 3.2, h - 1.4, items, size=7.5, spacing=1.1)


@slide_type("dept_category_questions")
def s_dept_category_questions(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    ins = ctx.ins() or {}
    _frame(slide, ctx, spec, f"{ctx.scope}｜カテゴリー・設問結果",
           f"{ctx.scope}の設問別平均を全庁と比較。全庁差が大きい設問を強み・課題の候補とする（左＝強み TOP、右＝課題 TOP）。")
    x, y, w, h = ctx.content()
    half = (w - 0.2) / 2
    for i, (title, kind, color) in enumerate((("強み（全庁比プラス TOP5）", "全庁比プラス", th.hexv("positive")),
                                              ("課題（全庁比マイナス TOP5）", "全庁比マイナス", th.hexv("negative")))):
        d = _ranking_rows(ctx, kind, 5)
        bx = x + i * (half + 0.2)
        add_text(slide, th, bx, y, half, 0.25, title, size=9, bold=True, color=color)
        cols = [{"id": "q", "title": "設問", "width": half - 1.9, "kind": "text", "chars": 24, "lines": 2},
                {"id": "mean", "title": "平均", "width": 0.55, "kind": "num"},
                {"id": "all", "title": "全庁", "width": 0.55, "kind": "num"},
                {"id": "gap", "title": "全庁差", "width": 0.8, "kind": "sig"}]
        rows = []
        for _, r in d.iterrows():
            q = r["question_id"]
            sc = ctx.row("設問", q)
            rows.append({"cells": {"q": f"{q} {ctx.qtext(q)}", "mean": r["mean"], "all": ctx.val("設問", q, "mean", "全庁", "全庁"),
                                   "gap": _sig_tuple(r["diff_all"], sc["sig_all"])}})
        p = ch.render_table(th, ctx.chart(f"dept_rank_{i}"), cols, rows or [{"cells": {"q": "該当なし"}}], row_h=0.42, header_h=0.36, font=7.2)
        place_image(slide, th, p, bx, y + 0.27, half, 2.5, align="left")
    # 下: カテゴリ表
    cats = ctx.summary["categories"]
    rows = [["カテゴリ", ctx.scope, "全庁", "全庁差", "前年差", "肯定率", "否定率"]]
    fills = {}
    for i, c in enumerate(cats, start=1):
        r = ctx.row("カテゴリ", c)
        rows.append([c, fmt(r["mean"]), fmt(ctx.val("カテゴリ", c, "mean", "全庁", "全庁")), fmt(r["diff_all"], signed=True),
                     fmt(r["diff_yoy"], signed=True) + _ref_mark(r["yoy_status"]), fmt_pct(r["pos"]), fmt_pct(r["neg"])])
        if r["sig_all"]:
            fills[(i, 3)] = th.hexv("positive_soft") if r["sig_all"] > 0 else th.hexv("negative_soft")
    rh = min(0.2, (h - 3.0) / len(rows))
    half_rows = (len(rows) + 1) // 2
    add_table(slide, th, x, y + 2.95, w, rows, col_widths=[2.6, 0.9, 0.9, 0.9, 0.9, 0.9, 0.9], row_h=rh, size=6.5, fills=fills) if len(rows) <= 8 else None
    if len(rows) > 8:
        left_rows = [rows[0]] + rows[1:half_rows]
        right_rows = [rows[0]] + rows[half_rows:]
        add_table(slide, th, x, y + 2.95, half, left_rows, col_widths=[1.6, 0.5, 0.5, 0.55, 0.6, 0.5, 0.5], row_h=0.2, size=6.3, fills={k: v for k, v in fills.items() if k[0] < half_rows})
        add_table(slide, th, x + half + 0.2, y + 2.95, half, right_rows, col_widths=[1.6, 0.5, 0.5, 0.55, 0.6, 0.5, 0.5], row_h=0.2, size=6.3,
                  fills={(k[0] - half_rows + 1, k[1]): v for k, v in fills.items() if k[0] >= half_rows})


@slide_type("dept_trend")
def s_dept_trend(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    if not ctx.has_prev:
        _frame(slide, ctx, spec, f"{ctx.scope}｜経年変化", "前年度データが無いため経年比較は行っていない。")
        return
    ov = ctx.row("指標", ctx.target)
    up, dn = _ranking_rows(ctx, "改善", 5), _ranking_rows(ctx, "悪化", 5)
    lead = f"{ctx.scope}の{ctx.target_name}は {fmt(ov['prev_mean'])} → {fmt(ov['mean'])}（前年差 {fmt(ov['diff_yoy'], signed=True)}{_ref_mark(ov['yoy_status'])}{'・有意' if ov['sig_yoy'] else ''}）。"
    if ov["prev_n"] and ov["prev_n"] < ctx.min_n:
        lead = f"{ctx.scope}の前年度回答者が {ctx.min_n} 名未満のため前年差は非表示。"
    _frame(slide, ctx, spec, f"{ctx.scope}｜経年変化", lead)
    x, y, w, h = ctx.content()
    half = (w - 0.2) / 2
    for i, (title, d, color) in enumerate((("改善 TOP5", up, th.hexv("positive")), ("悪化 BOTTOM5", dn, th.hexv("negative")))):
        bx = x + i * (half + 0.2)
        add_text(slide, th, bx, y, half, 0.25, title, size=9, bold=True, color=color)
        p = _ranking_table(ctx, ctx.chart(f"dept_trend_{i}"), d, half, with_dist=False)
        place_image(slide, th, p, bx, y + 0.27, half, 2.1, align="left")
    cats = ctx.summary["categories"]
    rows = ctx.rows("カテゴリ").set_index("unit").reindex(cats)
    mat = np.array([rows["diff_yoy"].values.astype(float)])
    p = ch.heatmap(th, ctx.chart("dept_yoy_heat"), [ctx.scope], [short_cat(c) for c in cats], mat, vlim=0.3, w=w, h=1.4, font=6.8,
                   col_chars=6, cbar_label="前年差", first_col_values=np.array([ov["mean"]]), first_col_title="総合")
    add_text(slide, th, x, y + 2.5, w, 0.25, "カテゴリ別 前年差", size=9, bold=True, color=th.hexv("brand_dark"))
    place_image(slide, th, p, x, y + 2.75, w, h - 2.8, align="left")


@slide_type("dept_segments")
def s_dept_segments(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    seg = _seg_rows(ctx, ctx.scope_type, ctx.scope)
    allseg = _seg_rows(ctx, "全庁", "全庁")
    lab = ctx.cfg.raw["segmentation"]["labels"]
    lead = (f"{ctx.scope}の職員タイプ構成。{lab['low_low']} {fmt_pct(seg.loc['low_low', 'share'])}（全庁 {fmt_pct(allseg.loc['low_low', 'share'])}）、"
            f"{lab['high_high']} {fmt_pct(seg.loc['high_high', 'share'])}（全庁 {fmt_pct(allseg.loc['high_high', 'share'])}）。")
    _frame(slide, ctx, spec, f"{ctx.scope}｜職員タイプ（4類型）", lead,
           footnote=ctx.footnote_std("　課別ポジショニングは課単位の所属データが無いため省略。"))
    x, y, w, h = ctx.content()
    _quadrant_diagram(slide, ctx, x, y, 5.4, h, seg, allseg, "全庁")
    labels = [ctx.scope, "全庁"] + ([f"{ctx.scope}（前年）"] if ctx.has_prev and seg["prev_share"].notna().any() else [])
    shares = [seg["share"].tolist(), allseg["share"].tolist()] + ([seg["prev_share"].tolist()] if len(labels) == 3 else [])
    p = ch.stacked100(th, ctx.chart("dept_seg"), labels, shares, [lab[k] for k in QUADRANTS], [th.seg_color(k) for k in QUADRANTS],
                      w=3.5, h=2.0, font=7, label_chars=9, show_pct_min=8)
    add_text(slide, th, x + 5.6, y, 3.5, 0.25, "全庁・前年との比較", size=9, bold=True, color=th.hexv("brand_dark"))
    place_image(slide, th, p, x + 5.6, y + 0.27, 3.5, 2.0, align="left")
    sd = ctx.res["segment_drivers"].head(5)
    items = [[{"text": f"{r['question_id']} ", "bold": True}, {"text": r["text"]}] for _, r in sd.iterrows()]
    add_text(slide, th, x + 5.6, y + 2.4, 3.5, 0.25, "タイプを分けるドライバー（全庁分析）", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x + 5.6, y + 2.67, 3.5, h - 2.7, items, size=7.2, marker="", spacing=1.05)


@slide_type("dept_key_driver")
def s_dept_key_driver(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    d, src = _kd(ctx)
    own = src == ctx.scope_type
    top = d.sort_values("impact", ascending=False).head(10)
    a = d[d["zone"] == "A"]
    lead = ((f"{ctx.scope}内での{ctx.target_name}との相関に基づく Key Driver（n={int(d['n_target'].iloc[0])}、相関のみ）。"
             if own else f"{ctx.scope}は n<{ctx.cfg.analysis['key_driver']['dept_min_n']} のため全庁の Key Driver を参照。")
            + f"A（最優先改善）ゾーン {len(a)} 問。")
    _frame(slide, ctx, spec, f"{ctx.scope}｜Key Driver", lead, footnote=_kd_footnote(ctx, src))
    x, y, w, h = ctx.content()
    labels = [f"{r['question_id']} {ctx.qtext(r['question_id'])}" for _, r in top.iterrows()]
    p = ch.hbar_scores(th, ctx.chart("dept_kd"), labels, top["r"].tolist(), xlim=(0, 0.8), w=5.6, h=h - 0.25,
                       colors=[th.zone_color(z) for z in top["zone"]], font=6.8, label_chars=18, left=0.44)
    add_text(slide, th, x, y - 0.02, 5.6, 0.22, "影響度上位10の相関係数（色＝ゾーン）", size=8.5, bold=True, color=th.hexv("brand_dark"))
    place_image(slide, th, p, x, y + 0.22, 5.7, h - 0.25, align="left")
    items = []
    for _, r in a.sort_values("impact", ascending=False).head(6).iterrows():
        q = r["question_id"]
        gap = ctx.val("設問", q, "diff_all")
        items.append([{"text": f"{q} ", "bold": True}, {"text": f"{ctx.qtext(q)}", "size": 7.5},
                      {"text": f"　r={fmt(r['r'])} / {ctx.scope} {fmt(ctx.val('設問', q))}（全庁差 {fmt(gap, signed=True)}）", "size": 6.8, "color": th.hexv("ink_soft")}])
    add_text(slide, th, x + 5.9, y, 3.2, 0.25, "A 最優先改善（高影響 × 低スコア）", size=9, bold=True, color=th.hexv("negative"))
    bullets(slide, th, x + 5.9, y + 0.27, 3.2, h - 0.3, items or ["該当なし"], size=7.5, marker="", spacing=1.05)


@slide_type("dept_actions")
def s_dept_actions(slide, ctx: Ctx, spec: dict):
    th = ctx.theme
    lib = ctx.cfg.action_library.get("themes", {})
    themes = _issue_themes(ctx, 3)
    ins = ctx.ins() or {}
    sen = ins.get("sentences", {})
    _frame(slide, ctx, spec, f"{ctx.scope}｜アクション仮説", sen.get("hint") or "課題テーマごとに部局で検討すべきテーマと、全庁施策との役割分担を整理する。",
           footnote="施策は検討候補。管理職WSで現場要因を確認したうえで、部局で扱う課題と人事・全庁で扱う課題を切り分ける。")
    x, y, w, h = ctx.content()
    cols = [{"id": "theme", "title": "課題テーマ", "width": 1.3, "kind": "text", "chars": 8, "lines": 2},
            {"id": "ev", "title": "根拠（全庁比マイナスの設問）", "width": 2.5, "kind": "text", "chars": 19, "lines": 3},
            {"id": "check", "title": "現場で確認すべき要因", "width": 1.6, "kind": "text", "chars": 12, "lines": 3},
            {"id": "dept", "title": "部局で検討", "width": 1.85, "kind": "text", "chars": 14, "lines": 3},
            {"id": "all", "title": "全庁・人事へ依頼", "width": w - 7.25, "kind": "text", "chars": 14, "lines": 3}]
    rows = []
    for t, qs in themes:
        L = lib.get(t, {})
        lv = L.get("level", {})
        rows.append({"cells": {"theme": t, "ev": "\n".join(qs[:2]), "check": "・".join(L.get("check", [])),
                               "dept": f"部局: {lv.get('部局', '')}／課: {lv.get('課', '')}", "all": lv.get("全庁", "")}})
    p = ch.render_table(th, ctx.chart("dept_actions"), cols, rows or [{"cells": {"theme": "該当なし"}}], row_h=0.78, header_h=0.36, font=7)
    place_image(slide, th, p, x, y, w, 2.85, align="left")
    ws = ctx.cfg.action_library.get("ws_questions", [])
    add_text(slide, th, x, y + 2.95, w, 0.25, "管理職ワークショップで使う問い", size=9, bold=True, color=th.hexv("brand_dark"))
    bullets(slide, th, x, y + 3.2, w, h - 3.2, [f"{i + 1}. {q}" for i, q in enumerate(ws)], size=7.5, marker="", spacing=1.05)
