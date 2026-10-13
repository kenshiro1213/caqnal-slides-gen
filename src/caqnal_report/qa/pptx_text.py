"""PPTX の中身を検査用に取り出す（本文・表・ノート・グラフのキャッシュ・埋込 Excel・非表示頁）。"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

NUM = re.compile(r"^-?\d+(?:\.\d+)?$")


@dataclass
class ChartData:
    slide: int
    title: str
    categories: list[str]
    values: list[float]              # 全系列のキャッシュ値（空は除く）
    workbook_values: list[float]     # 埋込 Excel の数値セル
    workbook_strings: list[str]


@dataclass
class SlideContent:
    no: int
    title: str
    hidden: bool
    texts: list[str] = field(default_factory=list)          # 図形・表の全文字列
    tables: list[list[list[str]]] = field(default_factory=list)
    notes: str = ""
    n_pictures: int = 0


@dataclass
class DeckContent:
    path: Path
    slides: list[SlideContent]
    charts: list[ChartData]

    def all_text(self) -> str:
        parts = []
        for s in self.slides:
            parts += s.texts + [s.notes]
        for c in self.charts:
            parts += c.categories + c.workbook_strings
        return "\n".join(parts)


def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _workbook(chart) -> tuple[list[float], list[str]]:
    from openpyxl import load_workbook

    part = chart.part.chart_workbook.xlsx_part
    if part is None:
        return [], []
    wb = load_workbook(io.BytesIO(part.blob), read_only=True, data_only=True)
    nums, strs = [], []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, bool) or v is None:
                    continue
                if isinstance(v, (int, float)):
                    nums.append(float(v))
                else:
                    strs.append(str(v))
    return nums, strs


def read_deck(path: Path) -> DeckContent:
    prs = Presentation(str(path))
    slides, charts = [], []
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        sc = SlideContent(no=i, title="", hidden=slide._element.get("show") == "0")  # noqa: SLF001
        for sh in _walk(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                sc.n_pictures += 1
            if sh.has_text_frame and sh.text_frame.text.strip():
                sc.texts.append(sh.text_frame.text)
                if not title and sh.top is not None and sh.top < 914400 * 1.0 and len(sh.text_frame.text) > 1:
                    title = sh.text_frame.text
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text for c in r.cells] for r in sh.table.rows]
                sc.tables.append(rows)
                sc.texts += [c for r in rows for c in r]
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                vals, cats = [], []
                for plot in ch.plots:
                    try:
                        cats += [str(c) for c in plot.categories]
                    except Exception:  # noqa: BLE001  散布図は categories を持たない
                        pass
                    for s in plot.series:
                        vals += [float(v) for v in s.values if v is not None]
                wn, ws = _workbook(ch)
                charts.append(ChartData(i, "", cats, vals, wn, ws))
        if slide.has_notes_slide:
            sc.notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
        sc.title = title
        for c in charts:
            if c.slide == i:
                c.title = title
        slides.append(sc)
    return DeckContent(path, slides, charts)
