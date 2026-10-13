"""verify 工程（仕様 v2 §16、プロンプト集 v2.1 の 8）。正常な run の検査と、わざと壊した PPTX の検出。"""
from __future__ import annotations

import copy

import pytest
from pptx import Presentation
from pptx.util import Inches

from caqnal_report.ingest.normalize import read_normalized
from caqnal_report.qa.verify import read_signoff, reader_decks, verify, write_signoff
from caqnal_report.run import Run
from caqnal_report.steps import run_until, step_verify


@pytest.fixture(scope="module")
def run(cfg, tmp_path_factory):
    c = copy.deepcopy(cfg)
    c.raw["output_dir"] = str(tmp_path_factory.mktemp("out"))
    c.raw.update({"publication": {"min_n": 5}, "aggregate": {"category_missing": "all_answered"},
                  "common9": {"status": "not_conducted"},
                  "models": {"cv": {"group_level": "部局", "folds": 5}, "tree": {"max_depth": 2, "min_leaf": 40},
                             "bootstrap": {"n": 10, "seed": 0}},
                  "readers": [{"id": "全庁", "kind": "all_staff"}, {"id": "総務部長", "kind": "dept_head", "unit": "U003"}]})
    c.raw["org_dictionary"] = "org.yaml"
    c.org_dictionary = {"levels": ["部局"], "units": [{"id": f"U{i:03d}", "name": n, "parent": None, "level": "部局",
                                                      "order": i} for i, n in enumerate(c.org_units, 1)]}
    r = Run.create(c, "qa")
    assert run_until(r, "render")
    return r


def _by_id(rep, prefix):
    return [c for c in rep.checks if c.id.startswith(prefix)]


def test_clean_run_has_no_ng(run):
    rep = verify(run)
    ng = [f"{c.id}: {c.detail} {c.items[:3]}" for c in rep.checks if c.status == "ng"]
    assert not ng, ng
    assert not rep.release_ready                                 # 人の確認が無いうちは納品版にしない
    assert all(c.status == "manual" for c in _by_id(rep, "人の確認"))
    scope = [c for c in _by_id(rep, "読者の範囲外") if "総務部長" in c.id]
    assert scope and all("範囲外 0" not in c.detail for c in scope)   # 部局長版は他部局を点検対象にしている
    assert [c for c in _by_id(rep, "記入シート 1 枚") if c.status == "ok"]


def test_verify_step_records_pending_status(run):
    assert step_verify(run)
    assert run.status == "draft_qa_pending"
    assert (run.step_dir("verify") / "qa_report.md").exists()


def test_broken_deck_is_detected(run):
    path = next(p for _, p in reader_decks(run)["全庁"] if "詳細付録" in p.name)
    original = path.read_bytes()
    comment = next(t for t in read_normalized(run.step_dir("normalize"))["comments"]["text"].dropna() if len(t) >= 12)
    try:
        prs = Presentation(str(path))
        for slide in prs.slides:                                   # 非表示の所属の行に数値を入れる
            for sh in slide.shapes:
                if sh.has_table:
                    for row in sh.table.rows:
                        if row.cells[0].text.strip() == "監査委員事務局":
                            row.cells[len(row.cells) - 1].text = "42"
        first = prs.slides[0]
        first.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = comment
        first.notes_slide.notes_text_frame.text = "八戸市 との比較"   # 他自治体名（ノート）
        prs.slides[1]._element.set("show", "0")                    # noqa: SLF001
        for slide in prs.slides:                                   # 表の点数を書き換える
            for sh in slide.shapes:
                if sh.has_table and len(sh.table.rows) > 1 and "所属" not in sh.table.cell(0, 0).text:
                    sh.table.cell(1, 1).text = "12.345"
                    break
        prs.save(str(path))
        rep = verify(run)
        ng = {c.id.split("（")[0] for c in rep.checks if c.status == "ng"}
        assert {"非表示の所属の行に数値が無い", "自由記述の原文が無い", "非表示頁が無い", "表の数値⇄公開データ",
                "他自治体の名称が無い", "同じ公開データから作り直すとバイト一致"} <= ng
    finally:
        path.write_bytes(original)


def test_signoff_is_tied_to_fingerprint(run):
    write_signoff(run, "確認者A", ["interpretation", "sources", "pages", "editable"], "2026-10-13")
    rep = verify(run)
    assert all(c.status == "ok" for c in _by_id(rep, "人の確認"))
    path = reader_decks(run)["全庁"][0][1]
    original = path.read_bytes()
    try:
        prs = Presentation(str(path))
        prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1)).text_frame.text = "修正"
        prs.save(str(path))
        rep = verify(run)                                          # 確認後に成果物が変われば確認は無効
        assert all(c.status == "manual" for c in _by_id(rep, "人の確認"))
        assert _by_id(rep, "確認記録と成果物の指紋")[0].status == "open"
        assert not rep.release_ready
    finally:
        path.write_bytes(original)
    assert read_signoff(run)["reviewer"] == "確認者A"
    with pytest.raises(ValueError):
        write_signoff(run, "確認者A", ["全部見た"])
