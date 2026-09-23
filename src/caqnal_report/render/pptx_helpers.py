"""python-pptx の薄いラッパー。テンプレート展開、レイアウト取得、テキスト・表・画像の絶対配置。

- 日本語フォントは `a:latin` と `a:ea` の両方に typeface を入れる（片方だけだとテーマフォントに戻る）
- 座標は全て inch
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .theme import Theme

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def open_template(theme: Theme) -> Presentation:
    prs = Presentation(str(theme.template))
    sld_ids = prs.slides._sldIdLst  # noqa: SLF001
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)
    return prs


def get_layout(prs: Presentation, theme: Theme, key: str):
    spec = theme.t["layouts"][key]
    if isinstance(spec, dict):
        name, midx = spec["name"], spec.get("master", theme.t["master_index"])
    else:
        name, midx = spec, theme.t["master_index"]
    master = prs.slide_masters[midx]
    for lo in master.slide_layouts:
        if lo.name == name:
            return lo
    raise KeyError(f"レイアウト {name} がマスター {midx} にありません")


def new_slide(prs: Presentation, theme: Theme, key: str = "body"):
    return prs.slides.add_slide(get_layout(prs, theme, key))


def set_run_font(run, theme: Theme, size: float | None = None, bold: bool | None = None,
                 color: str | None = None, italic: bool | None = None) -> None:
    f = run.font
    f.name = theme.family
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color.lstrip("#"))
    rPr = run._r.get_or_add_rPr()  # noqa: SLF001
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", theme.family)


def _fill_placeholder(slide, idx_candidates, text, theme, size=None, bold=None, color=None, align=None):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx in idx_candidates:
            tf = ph.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            if align:
                p.alignment = ALIGN[align]
            run = p.add_run()
            run.text = text
            set_run_font(run, theme, size, bold, color)
            return ph
    return None


def set_title(slide, theme: Theme, text: str) -> None:
    ph = _fill_placeholder(slide, (0,), text, theme, theme.sizes["title"], True, theme.hexv("ink"))
    if ph is None:
        add_text(slide, theme, *theme.box("title"), text, size=theme.sizes["title"], bold=True)


def set_subtitle(slide, theme: Theme, text: str, size: float | None = None) -> None:
    ph = _fill_placeholder(slide, (1,), text, theme, size or theme.sizes["lead"], False, theme.hexv("ink"))
    if ph is None:
        add_text(slide, theme, 0.55, 2.5, 8.0, 0.9, text, size=size or theme.sizes["lead"])


def set_page_number(slide, theme: Theme, n: int) -> None:
    filled = False
    for ph in slide.placeholders:
        if ph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(ph.placeholder_format.type):
            tf = ph.text_frame
            tf.clear()
            run = tf.paragraphs[0].add_run()
            run.text = str(n)
            set_run_font(run, theme, theme.sizes["small"], False, theme.hexv("ink_soft"))
            filled = True
    if not filled:
        x, y, w, h = theme.box("page_number")
        add_text(slide, theme, x, y, w, h, str(n), size=theme.sizes["small"], color=theme.hexv("ink_soft"),
                 align="right")


def set_lead(slide, theme: Theme, text: str) -> None:
    x, y, w, h = theme.box("lead")
    add_text(slide, theme, x, y, w, h, text, size=theme.sizes["lead"], color=theme.hexv("ink"), anchor="top")


def add_footnote(slide, theme: Theme, text: str) -> None:
    x, y, w, h = theme.box("footnote")
    add_text(slide, theme, x, y, w, h, text, size=theme.sizes["small"] - 0.5, color=theme.hexv("ink_soft"),
             anchor="bottom")


def add_text(slide, theme: Theme, x, y, w, h, text, size=None, bold=False, color=None, align="left",
             anchor="top", fill=None, line=None, margin=0.04, wrap=True, line_spacing=None):
    """text は str / list[str]（段落）/ list[list[dict]]（段落ごとの run 指定 {text,bold,color,size}）"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = ANCHOR[anchor]
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if fill:
        tb.fill.solid()
        tb.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    if line:
        tb.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
        tb.line.width = Pt(0.75)
    paras = text if isinstance(text, list) else [text]
    size = size or theme.sizes["body"]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if line_spacing:
            p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [{"text": para}]
        for r in runs:
            run = p.add_run()
            run.text = str(r.get("text", ""))
            set_run_font(run, theme, r.get("size", size), r.get("bold", bold), r.get("color", color or theme.hexv("ink")))
    return tb


def _strip_style(shape) -> None:
    """図形からテーマ参照（<p:style>）を外す。

    python-pptx の add_shape / add_connector はテーマの effectRef を持つ <p:style> を付けるため、
    shadow.inherit=False（空の effectLst）にしても LibreOffice 等ではテーマ既定の影が描かれる。
    塗り・線は明示的に設定しているので <p:style> ごと削除して影を完全に無くす。"""
    style = shape._element.find(qn("p:style"))  # noqa: SLF001
    if style is not None:
        shape._element.remove(style)  # noqa: SLF001
    shape.shadow.inherit = False


def add_rect(slide, theme: Theme, x, y, w, h, fill: str, line: str | None = None, shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    if line:
        shp.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    _strip_style(shp)
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return shp


def add_line(slide, theme: Theme, x1, y1, x2, y2, color: str, width_pt: float = 0.75):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = RGBColor.from_string(color.lstrip("#"))
    ln.line.width = Pt(width_pt)
    _strip_style(ln)
    return ln


def place_image(slide, theme: Theme, path: Path, x, y, max_w, max_h, align: str = "center", valign: str = "top"):
    with Image.open(path) as im:
        pw, ph = im.size
    w_in, h_in = pw / theme.dpi, ph / theme.dpi
    scale = min(max_w / w_in, max_h / h_in, 1.0)
    w, h = w_in * scale, h_in * scale
    dx = (max_w - w) / 2 if align == "center" else (max_w - w) if align == "right" else 0
    dy = (max_h - h) / 2 if valign == "middle" else 0
    return slide.shapes.add_picture(str(path), Inches(x + dx), Inches(y + dy), Inches(w), Inches(h))


def add_table(slide, theme: Theme, x, y, w, rows: list[list], col_widths: list[float] | None = None,
              row_h: float = 0.26, header: bool = True, size: float | None = None, header_fill: str | None = None,
              header_color: str | None = None, aligns: list[str] | None = None, fills: dict | None = None,
              colors: dict | None = None, bold_cells: set | None = None, zebra: bool = False):
    """rows[0] をヘッダーとする単純表。fills/colors は {(r,c): hex}。"""
    n_r, n_c = len(rows), len(rows[0])
    h = row_h * n_r
    shape = slide.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr  # noqa: SLF001
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Inches(w * cw / total)
    size = size or theme.sizes["small"]
    header_fill = header_fill or theme.hexv("brand_pale")
    for r in range(n_r):
        tbl.rows[r].height = Inches(row_h)
        for c in range(n_c):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(36000)
            cell.margin_top = cell.margin_bottom = Emu(18000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = rows[r][c]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = ALIGN[(aligns[c] if aligns else ("left" if c == 0 else "center"))]
            run = p.add_run()
            run.text = "" if val is None else str(val)
            is_head = header and r == 0
            color = (colors or {}).get((r, c)) or (header_color or theme.hexv("ink"))
            set_run_font(run, theme, size, is_head or ((r, c) in (bold_cells or set())), color)
            fill = (fills or {}).get((r, c))
            if is_head:
                fill = fill or header_fill
            elif zebra and r % 2 == 0:
                fill = fill or "F7F7F7"
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
            else:
                cell.fill.background()
    return shape


def kpi_card(slide, theme: Theme, x, y, w, h, title: str, value: str, sub: str = "", accent: str | None = None,
             value_size: float | None = None):
    accent = accent or theme.hexv("brand")
    add_rect(slide, theme, x, y, w, h, theme.hexv("bg"), theme.hexv("line"))
    add_rect(slide, theme, x, y, 0.06, h, accent)
    add_text(slide, theme, x + 0.12, y + 0.05, w - 0.2, 0.3, title, size=theme.sizes["small"] + 0.5,
             color=theme.hexv("ink_soft"))
    add_text(slide, theme, x + 0.12, y + 0.28, w - 0.2, h - 0.6, value, size=value_size or theme.sizes["kpi"],
             bold=True, color=theme.hexv("ink"), anchor="middle")
    if sub:
        add_text(slide, theme, x + 0.12, y + h - 0.34, w - 0.2, 0.3, sub, size=theme.sizes["small"],
                 color=theme.hexv("ink_soft"), anchor="bottom")


def bullets(slide, theme: Theme, x, y, w, h, items: list, size: float | None = None, marker: str = "■ ",
            color: str | None = None, marker_color: str | None = None, spacing: float = 1.1):
    paras = []
    for it in items:
        if isinstance(it, list):
            paras.append([{"text": marker, "color": marker_color or theme.hexv("brand")}] + it)
        else:
            paras.append([{"text": marker, "color": marker_color or theme.hexv("brand")}, {"text": str(it)}])
    return add_text(slide, theme, x, y, w, h, paras, size=size or theme.sizes["body"], color=color,
                    line_spacing=spacing)
