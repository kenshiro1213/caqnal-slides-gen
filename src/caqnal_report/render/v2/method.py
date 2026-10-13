"""社内向け 分析手法解説（仕様 v2 版）。

何をどう処理・分析したかを、その run の設定値・記録・公開データから説明する社内資料。
レポートと同じ run から作るので、数値と条件はレポートと一致する。
読むもの: run.json、各工程のログ（件数・設定のみ）、公開データ（全庁の読者）、非表示一覧、QA 記録、render の成果物。
個人単位の回答・自由記述の原文は読まない。社内資料なので点数の換算式も載せる（顧客向け本文には載せない）。

出力: output/runs/<run_id>/internal/<自治体>_<年度>_分析手法解説_社内向け.pptx
"""
from __future__ import annotations

import io
import json
from pathlib import Path

import pandas as pd
from pptx import Presentation

from ...config import ClientConfig
from ..pptx_helpers import add_line, add_rect
from .base import V2Theme, bar_chart, new_presentation, table, text
from .data import ALL, ReportData, f1, num
from .deck import _fixed_zip, deck_spec
from .slides import REGISTRY, Ctx, fill_toc, register

STEP_LABEL = {"validate": ("入力点検", "設定・原資料・組織辞書の点検。未確定を数える"),
              "normalize": ("標準化", "回答・名簿・設問・組織を標準の表へ。ID 照合・重複・除外を記録"),
              "aggregate": ("集計", "全庁・組織・属性 × 設問・カテゴリ・目的指標。相関"),
              "models": ("モデル", "重回帰（カテゴリ・設問）・決定木・群交差検証・採否候補"),
              "publish": ("公開データ", "一次・二次秘匿、読者別の範囲で切出し"),
              "compose": ("文章", "観測事実の下書き＋担当者の解釈（interpretation.yaml）"),
              "render": ("描画", "読者別の編集可能な PPTX"),
              "verify": ("品質検証", "数値照合・漏えい・内容・表示・再現性。人の確認で release_ready")}
STATUS_LABEL = [("draft_unconfirmed", "未確定あり"), ("validated", "入力確定"), ("computed", "分析済み"),
                ("public_view_ready", "公開データ済み"), ("draft_qa_pending", "検証待ち"), ("release_ready", "納品可")]
MODEL_LABEL = {"category": "カテゴリ重回帰", "question": "設問重回帰", "tree": "決定木"}


def _x(ctx: Ctx) -> dict:
    return ctx.extra


def _chap(spec: dict, default: str) -> str:
    return spec.get("chapter") or default


# ================================================================= 1. 全体像
@register("m_pipeline")
def m_pipeline(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    run = _x(ctx)["run"]
    p = ctx.page("8 工程の流れと今回の実行結果", "工程はファイルだけで受け渡し、前の工程が通らなければ先へ進まない",
                 chapter=_chap(spec, "1. 全体像"),
                 note=f"run {run['run_id']}（{run['created']} 作成）。各工程の入力・設定・コードの指紋を run.json に記録。"
                      "private（個人単位）のデータは publish より後の工程に渡さない。")
    steps = list(STEP_LABEL)
    w, h, gap = (p.width - 0.3 * 3) / 4, 2.2, 0.3
    for i, s in enumerate(steps):
        x = p.left + (i % 4) * (w + gap)
        y = p.body_top + (i // 4) * (h + 0.45)
        rec = run["steps"].get(s) or {}
        state = rec.get("state", "未実行")
        color = {"passed": th.hexv("navy"), "failed": th.hexv("neg"), "blocked": th.hexv("neg")}.get(state, th.hexv("ink_soft"))
        add_rect(p.slide, th, x, y, w, h, th.hexv("bg"), th.hexv("line"))
        add_rect(p.slide, th, x, y, w, 0.5, color)
        name, desc = STEP_LABEL[s]
        text(p, x + 0.15, y, w - 0.3, 0.5, f"{i + 1}. {name}（{s}）", size=th.sizes["small"], bold=True,
             color="FFFFFF", anchor="middle")
        extra = []
        if rec.get("unresolved"):
            extra.append(f"未確定 {len(rec['unresolved'])}")
        if rec.get("warnings"):
            extra.append(f"警告 {len(rec['warnings'])}")
        text(p, x + 0.15, y + 0.6, w - 0.3, h - 0.7,
             [desc, f"結果: {state}" + (f"（{' / '.join(extra)}）" if extra else "")], size=th.sizes["note"] + 0.5)
        if i % 4 != 3:
            add_line(p.slide, th, x + w, y + h / 2, x + w + gap, y + h / 2, th.hexv("gold"), 1.5)


@register("m_run_status")
def m_run_status(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    run = _x(ctx)["run"]
    cur = run["status"]
    p = ctx.page("run の状態と再実行の規則", f"今回の run は「{dict(STATUS_LABEL)[cur]}（{cur}）」",
                 chapter=_chap(spec, "1. 全体像"),
                 note="unknown・failed・blocked を passed として扱わない。release_ready の run は変更できない。")
    w = (p.width - 0.2 * 5) / 6
    idx = [s for s, _ in STATUS_LABEL].index(cur)
    for i, (s, lbl) in enumerate(STATUS_LABEL):
        x = p.left + i * (w + 0.2)
        fill = th.hexv("navy") if i <= idx else th.hexv("pale")
        add_rect(p.slide, th, x, p.body_top, w, 0.9, fill, th.hexv("gold") if i == idx else None)
        text(p, x, p.body_top, w, 0.9, [lbl, s], size=th.sizes["note"], bold=i == idx,
             color="FFFFFF" if i <= idx else th.hexv("ink"), align="center", anchor="middle")
    rows = [["規則", "内容"],
            ["新しい run を作るとき", "client.yaml・設問・組織辞書・共通設定、原資料、コードのどれかが変わったとき（既存の成果は上書きしない）"],
            ["同じ run でやり直すとき", "その工程より後の記録と出力を消す（古い出力の取り違えを防ぐ）"],
            ["指紋", f"コード {run['code'][:12]}…／設定 {run['config'][:12]}…／原資料 {len(run['inputs'])} 件"],
            ["納品可にする条件", "verify で NG・未完が無く、人の確認（確認者・日付・確認項目）が今の成果物の指紋と一致"]]
    table(p, p.left, p.body_top + 1.3, p.width, rows, widths=[1.6, 6], row_h=0.6, aligns=["left", "left"],
          size=th.sizes["small"])


# ================================================================= 2. 入力と設定
@register("m_settings")
def m_settings(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    cfg = d.cfg
    m = cfg.models
    pub = cfg.raw.get("publication") or {}
    objs = "、".join(f"{o['question']} {o.get('name')}" for o in d.objectives)
    readers = cfg.readers or [{"id": ALL, "kind": "all_staff"}]
    kinds = pd.Series([r["kind"] for r in readers]).value_counts().to_dict()
    p = ctx.page("案件設定（client.yaml）の要点", "数字に影響する条件はすべて案件設定から読み、共通コードに固定値を置かない",
                 chapter=_chap(spec, "2. 入力と集計"),
                 note="null＝確認待ち。確認待ちが残る限り run は draft_unconfirmed から進めても納品可にならない。")
    rows = [["項目", "今回の値", "設定の場所"],
            ["最低公開人数", f"{pub.get('min_n')} 名（二次秘匿 {'あり' if pub.get('secondary_suppression') else 'なし'}）", "publication"],
            ["目的指標", objs, "objectives"],
            ["カテゴリ", f"{len(d.categories)} 個（欠損: {(cfg.raw.get('aggregate') or {}).get('category_missing')}）", "questions / aggregate"],
            ["説明設問", f"{len(d.drivers)} 問（条件付き設問はモデルの説明変数から除外）", "questions"],
            ["共通9問", (cfg.raw.get("common9") or {}).get("status"), "common9"],
            ["属性", "、".join(d.attributes()) or "なし", "attributes / attribute_rules"],
            ["組織", f"{len(d.tables['org'])} 組織（階層: {'・'.join((cfg.org_dictionary or {}).get('levels') or [cfg.org_level])}）", "org_dictionary"],
            ["読者", "、".join(f"{k} {v}" for k, v in kinds.items()), "readers"],
            ["検証", f"{(m.get('cv') or {}).get('group_level')} 単位 {(m.get('cv') or {}).get('folds')} 分割", "models.cv"],
            ["決定木・再抽出", f"深さ {(m.get('tree') or {}).get('max_depth')}・末端 {(m.get('tree') or {}).get('min_leaf')} 名・"
                          f"再抽出 {(m.get('bootstrap') or {}).get('n')} 回", "models.tree / bootstrap"],
            ["フォント", cfg.theme["font"]["presets"][cfg.theme["font"]["preset"]]["family"], "theme.font.preset"]]
    table(p, p.left, p.body_top, p.width, rows, widths=[1.4, 5, 1.6], row_h=0.42, aligns=["left", "left", "left"],
          size=th.sizes["small"])


@register("m_unresolved")
def m_unresolved(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    run = _x(ctx)["run"]
    items = [(s, u) for s, rec in run["steps"].items() if s != "verify" for u in rec.get("unresolved", [])]
    p = ctx.page("確認待ちの事項", f"{len(items)} 件。確定するまで該当する数値・頁は納品版にしない" if items else "確認待ちはない",
                 chapter=_chap(spec, "2. 入力と集計"),
                 note="確定したら client.yaml に値と確認日を記録する。他自治体の補正や値を自動で当てはめない。")
    rows = [["工程", "事項"]] + [[STEP_LABEL.get(s, (s,))[0], u] for s, u in items] or [["－", "なし"]]
    table(p, p.left, p.body_top, p.width, rows, widths=[1.2, 7], row_h=min(0.5, 6 / max(len(rows), 1)),
          aligns=["left", "left"], size=th.sizes["small"])


@register("m_normalize")
def m_normalize(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    log = _x(ctx)["logs"]["normalize"].get("log", {})
    stages = log.get("stages") or {}
    p = ctx.page("標準化と照合の結果", f"回答 {num(log.get('responses_rows'))} 行のうち {num(log.get('responses_used'))} 行を集計に使用",
                 chapter=_chap(spec, "2. 入力と集計"),
                 note="ID は文字列で照合（全角半角・空白・ゼロ埋めを規則で統一）。重複は規則（keep_last 等）が無ければ止める。"
                      "生年月日・入職日は基準日で区分にしてから保存し、元の値は残さない。")
    rows = [["段階", "件数"]] + [[k, "－" if v is None else num(v)] for k, v in stages.items()]
    table(p, p.left, p.body_top, p.width * 0.45, rows, widths=[3, 1.2], row_h=0.42, aligns=["left", "right"],
          size=th.sizes["small"])
    miss = log.get("missing") or {}
    rows2 = [["欠損の理由", "回答数", "扱い"]] + [[k, num(v), {"条件外": "対象外の職員。分母から除く（カテゴリでは回答不要）",
                                                         "未回答": "分母から除く", "除外選択肢": "「わからない」等。分母から除く"}.get(k, "")]
                                           for k, v in miss.items()]
    x = p.left + p.width * 0.5
    table(p, x, p.body_top, p.width * 0.5, rows2 if len(rows2) > 1 else rows2 + [["なし", "", ""]],
          widths=[1.2, 0.8, 3.2], row_h=0.42, aligns=["left", "right", "left"], size=th.sizes["small"])
    corr = log.get("corrections") or []
    text(p, x, p.body_top + 0.42 * (max(len(rows2), 2) + 1) + 0.2, p.width * 0.5, 1.2,
         f"設問対応の補正: {len(corr)} 件" + ("（理由と確認者を client.yaml に記録）" if corr else ""), size=th.sizes["small"])


@register("m_scoring")
def m_scoring(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    cfg = d.cfg
    o = d.objectives[0] if d.objectives else None
    st = d.stat(ALL, ALL, "目的指標", o["id"]) if o else None
    sname = (cfg.q(o["question"]).get("scale") if o else None) or next(iter(cfg.raw["scales"]))
    scale = cfg.scale(sname)
    lo, hi = scale.get("min"), scale.get("max")
    ex = ""
    if st:
        ex = (f"例: {o['question']} {o.get('name')} の平均 {st['mean']:.3f} → "
              f"({st['mean']:.3f} − {lo}) ÷ ({hi} − {lo}) × 100 ＝ {st['score']:.1f} 点")
    p = ctx.page("点数と肯定・中立・否定の計算", "点数＝平均を尺度の最小〜最大で 0〜100 に置き換えた値（社内向け。顧客向け本文には式を載せない）",
                 chapter=_chap(spec, "2. 入力と集計"),
                 note="丸めは表示のときだけ。集計値は丸め前で保存し、全庁差・相関などの計算にも丸め前を使う。")
    lines = [f"点数 ＝ (平均 − {lo}) ÷ ({hi} − {lo}) × 100", ex,
             "逆転設問は数値化のときに 1 回だけ反転し、肯定・否定は反転後の値で分類する。",
             "カテゴリの点数は、構成設問の平均を回答者ごとに出してから集計する（欠損規則は案件設定）。",
             "中立が無い尺度（4 件法など）では中立を作らない。"]
    text(p, p.left, p.body_top, p.width * 0.55, 3.6, [x for x in lines if x], size=th.sizes["body"], line_spacing=1.5)
    rows = [["選択肢", "値", "分類"]]
    for lbl, v in scale["labels"].items():
        cls = "肯定" if v in scale.get("positive", []) else "中立" if v in (scale.get("neutral") or []) else "否定"
        rows.append([lbl, str(v), cls])
    table(p, p.left + p.width * 0.6, p.body_top, p.width * 0.4, rows, widths=[2.5, 0.7, 0.9], row_h=0.42,
          aligns=["left", "center", "center"], size=th.sizes["small"])


@register("m_aggregate")
def m_aggregate(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    log = _x(ctx)["logs"]["aggregate"].get("log", {})
    st = d.tables["stats"]
    rec = log.get("reconcile") or {}
    p = ctx.page("集計の範囲と照合", f"集計対象 {num(log.get('analysed_n'))} 名。最上位組織の合計＋所属なし＝全庁を確認",
                 chapter=_chap(spec, "2. 入力と集計"),
                 note="組織の値は配下を含む。属性別の回答率は対象者リストを分母にする（名簿が無ければ計算不能と表示）。")
    rows = [["集計の単位", "数"],
            ["組織の範囲", num(st.loc[st.scope_type == "組織", "scope_id"].nunique())],
            ["属性の区分", num(st.loc[st.scope_type == "属性", "scope_id"].nunique())],
            ["設問", num(st.loc[st.unit_type == "設問", "unit_id"].nunique())],
            ["カテゴリ", num(st.loc[st.unit_type == "カテゴリ", "unit_id"].nunique())],
            ["目的指標", num(st.loc[st.unit_type == "目的指標", "unit_id"].nunique())],
            ["集計セル（行）", num(len(st))],
            ["相関（項目 × 目的指標）", num(log.get("n_correlations"))]]
    table(p, p.left, p.body_top, p.width * 0.45, rows, widths=[3, 1.2], row_h=0.42, aligns=["left", "right"],
          size=th.sizes["small"])
    rows2 = [["照合", "人数"]] + [[k, num(v)] for k, v in rec.items()]
    table(p, p.left + p.width * 0.5, p.body_top, p.width * 0.4, rows2, widths=[3, 1.2], row_h=0.42,
          aligns=["left", "right"], size=th.sizes["small"])


# ================================================================= 3. 匿名性と読者
@register("m_anonymity")
def m_anonymity(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    sup: pd.DataFrame = _x(ctx)["suppression"]
    scopes = sup[sup.unit_type.isna()]
    cells = sup[sup.unit_type.notna()]
    names = dict(zip(d.tables["org"].unit_id, d.tables["org"].name))
    p = ctx.page("匿名性: 一次秘匿と二次秘匿", f"組織・区分 {len(scopes)} 件、セル {len(cells)} 件を非表示",
                 chapter=_chap(spec, "3. 匿名性と読者"),
                 note="非表示の値は公開データに残さない（行は残して値を空にし、理由を記録）。render と verify は公開データだけを読む。")
    text(p, p.left, p.body_top, p.width * 0.45, 3.4, [
        f"一次: 回答者または有効 n が最低公開人数（{d.manifest.get('min_n')} 名）未満の組織・属性区分・セルを隠す。",
        "二次: 親（上位組織・全庁）から公開中の子を差し引いた残りが 1〜(最低人数−1) 名になるとき、"
        "公開中の子のうち最も少ないものを追加で隠す。回答者数とセルの n の両方で行う。",
        "決定木の葉・ノードも最低人数未満なら人数と平均を隠す。",
        "公開条件とモデルの採用条件は別。人数が足りても分析が安定するとは限らない。"], size=th.sizes["small"], line_spacing=1.5)
    rows = [["非表示の組織・区分", "理由"]]
    for r in scopes.itertuples():
        rows.append([names.get(r.scope_id, str(r.scope_id).replace("=", "：")), r.reason])
    if len(rows) == 1:
        rows.append(["なし", ""])
    table(p, p.left + p.width * 0.5, p.body_top, p.width * 0.5, rows[:14], widths=[2, 3], row_h=0.4,
          aligns=["left", "left"], size=th.sizes["note"] + 0.5)
    by = cells.groupby("unit_type").size().to_dict()
    text(p, p.left + p.width * 0.5, p.body_top + 0.4 * min(len(rows), 14) + 0.2, p.width * 0.5, 0.8,
         "非表示セルの内訳: " + ("、".join(f"{k} {v}" for k, v in by.items()) or "なし"), size=th.sizes["note"] + 0.5)


@register("m_readers")
def m_readers(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    decks = _x(ctx)["decks"]
    p = ctx.page("読者別の成果物と閲覧範囲", f"{len({r for r, _, _ in decks})} 読者に {len(decks)} ファイル",
                 chapter=_chap(spec, "3. 匿名性と読者"),
                 note="部局長・課長は同じ部局内だけ。部局が決まらなければ止める（全庁へ広げない）。自所属が非表示なら数値の無い短い版にする。")
    rows = [["読者", "ファイル", "頁"]] + [[r, f, num(n)] for r, f, n in decks]
    chunks = [rows[1:][i:i + 15] for i in range(0, len(rows) - 1, 15)] or [[]]
    w = p.width / len(chunks[:2]) - (0.2 if len(chunks) > 1 else 0)
    for k, ch in enumerate(chunks[:2]):
        table(p, p.left + k * (w + 0.4), p.body_top, w, [rows[0]] + ch, widths=[2.2, 4.2, 0.6], row_h=0.36,
              aligns=["left", "left", "right"], size=th.sizes["note"])


# ================================================================= 4. 分析
@register("m_kd")
def m_kd(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    p = ctx.page("キードライバー: 相関", "カテゴリ・説明設問と各目的指標の相関を、組合せごとの n と 95% 区間つきで出す",
                 chapter=_chap(spec, "4. 分析"),
                 note="Pearson の相関、区間は Fisher の z 変換。目的指標は説明変数に入れない。点数下位と相関上位をそのまま施策の順位にしない。")
    objs = d.objectives
    w = (p.width - 0.3 * (len(objs) - 1)) / max(len(objs), 1)
    for i, o in enumerate(objs):
        c = d.corr(o["id"], "カテゴリ").sort_values("r", ascending=False).head(6)
        rows = [[f"{o.get('short') or o['name']}", "r", "95%区間", "n"]] + [
            [r.unit_id, f"{r.r:.3f}", f"{r.ci_low:.2f}～{r.ci_high:.2f}", num(r.n)] for r in c.itertuples()]
        table(p, p.left + i * (w + 0.3), p.body_top, w, rows, widths=[2.6, 0.8, 1.4, 0.7], row_h=0.42,
              aligns=["left", "center", "center", "right"], size=th.sizes["note"] + 0.5)


@register("m_ols")
def m_ols(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    f = d.tables["model_fit"]
    f = f[f.model != "tree"]
    interval = d.cfg.models.get("ols", {}).get("interval") or "HC3"
    p = ctx.page("重回帰の診断", "推定できるか（rank・自由度）、共線性（VIF）、当てはまり、組織単位の感度を確認する",
                 chapter=_chap(spec, "4. 分析"),
                 note=f"標準化係数 β、区間は {interval}。感度分析＝{d.cfg.models.get('cv', {}).get('group_level')}単位のクラスタ頑健区間で、"
                      "主の区間では 0 を含まないのに 0 を含むようになる項目数。カテゴリと構成設問を同じモデルに入れない。")
    rows = [["目的指標", "モデル", "n", "変数", "rank", "残差自由度", "VIF 最大", "R²", "調整済み R²", "感度で不安定"]]
    for r in f.itertuples():
        rows.append([d.obj_name(r.objective_id, True), MODEL_LABEL[r.model], num(r.n), num(r.k), num(r.rank),
                     num(r.df_resid), f1(r.vif_max), f1(r.r2, 3), f1(r.adj_r2, 3),
                     num(getattr(r, "cluster_sign_unstable", None))])
    table(p, p.left, p.body_top, p.width, rows, widths=[1, 1.3, 0.7, 0.6, 0.6, 0.9, 0.8, 0.7, 0.9, 1],
          row_h=0.46, aligns=["left", "left"] + ["center"] * 8, size=th.sizes["small"])


@register("m_validation")
def m_validation(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    f = d.tables["model_fit"]
    cv = d.cfg.models.get("cv") or {}
    p = ctx.page("未学習データでの検証", f"{cv.get('group_level')}単位で {cv.get('folds')} 分割し、学習に使わなかった{cv.get('group_level')}の職員を予測する",
                 chapter=_chap(spec, "4. 分析"),
                 note="標準化・前処理は各分割の学習側だけで行う。基準は学習側の平均だけで予測した場合の MAE。"
                      "同じ組織の職員が学習と検証に分かれると、組織が共有する条件が漏れて当てはまりを過大に見積もる。")
    cats = [f"{d.obj_name(r.objective_id, True)}・{MODEL_LABEL[r.model]}" for r in f.itertuples()]
    bar_chart(p, p.left, p.body_top, p.width * 0.55, p.body_bottom - p.body_top, cats,
              {"検証 MAE": [round(v, 1) for v in f.cv_mae], "平均のみ MAE": [round(v, 1) for v in f.base_mae]},
              [th.hexv("navy"), th.hexv("line")], legend=True)
    rows = [["モデル", "検証 R²", "群数"]] + [[c, f1(r.cv_r2, 3), num(r.cv_groups)] for c, r in zip(cats, f.itertuples())]
    table(p, p.left + p.width * 0.6, p.body_top, p.width * 0.4, rows, widths=[3, 1, 0.7], row_h=0.42,
          aligns=["left", "center", "center"], size=th.sizes["small"])


@register("m_tree")
def m_tree(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    stab = d.tables.get("tree_stability", pd.DataFrame())
    t = d.cfg.models.get("tree") or {}
    b = d.cfg.models.get("bootstrap") or {}
    p = ctx.page("決定木と根元の安定性", f"深さ {t.get('max_depth')}・末端 {t.get('min_leaf')} 名。{b.get('n')} 回の再抽出で根元に選ばれる項目を数える",
                 chapter=_chap(spec, "4. 分析"),
                 note=f"再抽出は{d.cfg.models.get('cv', {}).get('group_level')}単位。根元の安定性は木全体の安定性ではない。"
                      "閾値はこの標本での条件で、配置・評価の基準値にしない。")
    objs = d.objectives
    w = (p.width - 0.3 * (len(objs) - 1)) / max(len(objs), 1)
    for i, o in enumerate(objs):
        s = stab[stab.objective_id == o["id"]].sort_values("root_freq", ascending=False).head(5) if not stab.empty else stab
        rows = [[f"{o.get('short') or o['name']}の根元", "選ばれた割合"]] + [
            [d.qlabel(r.feature, 18), f"{r.root_freq:.0%}"] for r in s.itertuples()]
        table(p, p.left + i * (w + 0.3), p.body_top, w, rows if len(rows) > 1 else rows + [["未計算", ""]],
              widths=[3.2, 1], row_h=0.45, aligns=["left", "center"], size=th.sizes["note"] + 0.5)


@register("m_adoption")
def m_adoption(ctx: Ctx, spec: dict) -> None:
    th, d = ctx.th, ctx.d
    f = d.tables["model_fit"]
    ref = d.cfg.analysis.get("models_v2", {})
    p = ctx.page("モデルの採否", "候補は規則で出し、最終判断は担当者が理由とともに models.decisions に記録する",
                 chapter=_chap(spec, "4. 分析"),
                 note=f"規則: 推定できない→推定不能／検証 MAE が平均のみ以上、または検証 R² ≤ 0→保留／VIF 最大 > {ref.get('vif_reference', 10)}、"
                      f"根元の最頻割合 < {ref.get('stability_reference', 0.8):.0%}、検証群が分割数未満→参考／それ以外→採用。保留・推定不能の係数は本文に載せない。")
    rows = [["目的指標", "モデル", "候補", "担当者の判断", "理由"]]
    for r in f.itertuples():
        rows.append([d.obj_name(r.objective_id, True), MODEL_LABEL[r.model], r.candidate,
                     r.decision if isinstance(r.decision, str) else "（未記録）",
                     r.reasons if isinstance(r.reasons, str) else ""])
    table(p, p.left, p.body_top, p.width, rows, widths=[1, 1.4, 0.8, 1.2, 5], row_h=0.46,
          aligns=["left", "left", "center", "center", "left"], size=th.sizes["small"])


# ================================================================= 5. 文章・品質
@register("m_compose")
def m_compose(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    qa = _x(ctx).get("qa")
    ph = []
    if qa:
        ph = [c for c in qa["checks"] if c["id"].startswith("未記入の枠") and c["status"] == "open"]
    p = ctx.page("文章の作り方と担当", "数値はツール、文章案は AI、判断と確定は担当者",
                 chapter=_chap(spec, "5. 文章と品質"),
                 note="interpretation.yaml は自治体から受け取るものではなく、CAQNAL の担当者が書く（AI の下書きを使ってよいが、確認して確定する）。")
    rows = [["誰が", "何をする", "どこに残す"],
            ["ツール", "集計・モデル・公開データ、観測事実の下書き（点数の低いカテゴリ・相関・最も低い設問）", "run/compose/<読者>/compose.json"],
            ["AI（下書き）", "示唆・背景仮説・方針・施策の文章案。観測事実への参照を付ける", "interpretation.yaml（下書き）"],
            ["担当者", "設問対応の補正、モデルの採否、自由記述の公開要約、施策の実現性、納品範囲を確認して確定", "interpretation.yaml・client.yaml"],
            ["担当者（確認）", "解釈・出典・全頁の目視・アプリでの編集性を確認", "caqnal signoff → signoff.json"]]
    table(p, p.left, p.body_top, p.width, rows, widths=[1.3, 5, 2.4], row_h=0.6, aligns=["left", "left", "left"],
          size=th.sizes["small"])
    text(p, p.left, p.body_top + 0.6 * len(rows) + 0.3, p.width, 0.8,
         f"今回の未記入: {sum(len(c['items']) for c in ph)} 頁（{len(ph)} ファイル）" if qa else "verify 未実行のため未記入の頁数は未集計",
         size=th.sizes["small"], color=th.hexv("ink_soft"))


@register("m_qa")
def m_qa(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    qa = _x(ctx).get("qa")
    if not qa:
        p = ctx.page("品質検証（verify）", "この run はまだ verify していない", chapter=_chap(spec, "5. 文章と品質"))
        text(p, p.left, p.body_top, p.width, 1, "caqnal verify --client <自治体> で検証する。", size=th.sizes["body"])
        return
    from ...qa.verify import AREAS

    checks = pd.DataFrame(qa["checks"])
    cnt = checks.groupby(["area", "status"]).size().unstack(fill_value=0)
    p = ctx.page("品質検証（verify）の結果", f"NG {int((checks.status == 'ng').sum())}・未完 {int((checks.status == 'open').sum())}・"
                 f"人の確認 {int((checks.status == 'manual').sum())}・OK {int((checks.status == 'ok').sum())}",
                 chapter=_chap(spec, "5. 文章と品質"),
                 note="自動で確認できないこと（解釈・出典・全頁の目視・アプリでの編集）は「人の確認」として残し、ツールは合格と書かない。")
    rows = [["監査区分", "OK", "NG", "未完", "人の確認"]]
    for a, name in AREAS.items():
        r = cnt.loc[a] if a in cnt.index else {}
        rows.append([f"{a}. {name}"] + [num(int(r.get(s, 0))) if len(r) else "0" for s in ("ok", "ng", "open", "manual")])
    table(p, p.left, p.body_top, p.width * 0.55, rows, widths=[2.4, 0.6, 0.6, 0.6, 0.9], row_h=0.44,
          aligns=["left"] + ["center"] * 4, size=th.sizes["small"])
    so = qa.get("signoff")
    text(p, p.left + p.width * 0.6, p.body_top, p.width * 0.4, 3, [
        f"判定: {'release_ready（納品可）' if qa.get('release_ready') else '納品不可（draft_qa_pending）'}",
        "確認者: " + (f"{so['reviewer']}（{so['date']}）" if so else "未記録"),
        "数値照合: private 集計⇄公開データ⇄PPT の表・グラフ・埋込 Excel",
        "漏えい: 非表示の行・範囲外の所属名・自由記述原文・他自治体名・非表示頁（ノート・埋込 Excel を含む）"],
         size=th.sizes["small"])


@register("m_limits")
def m_limits(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    p = ctx.page("分析の限界と書き方の注意", "相関・係数・分岐は原因や施策の効果を示さない", chapter=_chap(spec, "5. 文章と品質"))
    items = ["横断的な自己回答の調査。同じ調査への複数の分析は独立した証拠ではない。",
             "相関を原因、β を施策の効果、継続意向を離職率、推奨意向を NPS として説明しない。",
             "多くの項目を同時に探索している。属性の統制はしていない（必要なら追加で確認する）。",
             "所属別の数値は、少人数では偶然の振れが大きい。順位付けや課長個人の評価に使わない。",
             "調査だけで必要な増員人数・能力・格差・現行運用の程度を断定しない。",
             "前回比較は、設問・選択肢・母集団・集計定義が対応する場合だけ行う。"]
    text(p, p.left, p.body_top, p.width, 5, ["・" + i for i in items], size=th.sizes["body"], line_spacing=1.6)


@register("m_files")
def m_files(ctx: Ctx, spec: dict) -> None:
    th = ctx.th
    p = ctx.page("設定ファイルと出力の場所", "自治体ごとの違いは設定に、処理は共通コードに置く", chapter=_chap(spec, "5. 文章と品質"),
                 note="再実行: uv run caqnal report --client <自治体> → caqnal verify → 確認後 caqnal signoff。社内向け資料は caqnal method。")
    rows = [["ファイル", "役割"],
            ["local-governments/<自治体>/client.yaml", "入力・公開条件・目的指標・尺度・属性・読者・分析条件・フォント"],
            ["local-governments/<自治体>/questions.yaml", "設問全文・カテゴリ・逆転・条件"],
            ["local-governments/<自治体>/org.yaml", "組織辞書（ID・親・名称・階層・表示順）"],
            ["local-governments/<自治体>/interpretation.yaml", "担当者の解釈・方針・施策・確認指標・職員の声の要約"],
            ["config/decks_v2/*.yaml", "読者別デッキの頁構成"],
            ["config/theme.yaml・analysis_defaults.yaml", "色・フォント・採否の基準値"],
            ["output/runs/<run_id>/", "run.json、private/（個人単位）、public/<読者>/、compose/、render/、qa/、internal/"],
            ["output/release/<run_id>/", "納品版（verify と人の確認が揃った版。上書きしない）"]]
    table(p, p.left, p.body_top, p.width, rows, widths=[3.6, 5], row_h=0.5, aligns=["left", "left"],
          size=th.sizes["small"])


# ================================================================= 組立
def _deck_list(run_dir: Path) -> list[tuple[str, str, int]]:
    out = []
    for p in sorted((run_dir / "render").glob("*.pptx")):
        if p.name.startswith("~$"):              # PowerPoint で開いているときのロックファイル
            continue
        n = len(Presentation(str(p)).slides)
        reader = p.stem
        out.append((reader, p.name, n))
    return out


def build_method(cfg: ClientConfig, run) -> Path:
    from ...analysis.publish import read_public

    def _json(path: Path) -> dict:
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}

    readers = {r["id"]: r for r in (cfg.readers or [{"id": ALL, "kind": "all_staff"}])}
    rid = next((k for k, r in readers.items() if r["kind"] in ("all_staff", "hr")), next(iter(readers)))
    data = ReportData.load(cfg, run.step_dir("publish"), run.step_dir("compose"), rid)
    manifests = {r: read_public(run.step_dir("publish"), r)[1] for r in readers}
    decks = []
    for name, file, n in _deck_list(run.dir):
        who = next((r for r in sorted(readers, key=len, reverse=True) if f"_{r}" in file or f"{r}向け" in file), ALL)
        decks.append((who, file, n))
    extra = {
        "run": run.data,
        "logs": {s: _json(run.step_dir(s) / f"{s}_log.json") for s in ("normalize", "aggregate", "models")},
        "suppression": pd.read_csv(run.step_dir("publish") / "suppression.csv", dtype={"scope_id": str, "unit_id": str}),
        "qa": _json(run.step_dir("verify") / "qa_report.json") or None,
        "decks": decks,
        "manifests": manifests,
    }
    th = V2Theme(cfg)
    prs = new_presentation(th)
    c = cfg.client
    ctx = Ctx(prs=prs, th=th, d=data, label=f"{c['name']}｜社内向け", kind="分析手法解説（社内向け）",
              footer=f"© CAQNAL CO.,LTD.  •  {c['name']} {c.get('fiscal_year_label', '')} 分析手法解説（社内向け・外部配布不可）",
              extra=extra)
    deck = deck_spec(cfg, "手法解説")
    for spec in deck["slides"]:
        REGISTRY[spec["type"]](ctx, spec)
    fill_toc(ctx)
    out = run.dir / "internal" / deck["filename"].format(client=c["name"], fy=cfg.fiscal_year)
    out.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    prs.save(buf)
    out.write_bytes(_fixed_zip(buf.getvalue()))
    return out
