import hashlib
from pathlib import Path

from pptx import Presentation

from caqnal_report.analysis.pipeline import load_intermediate, write_intermediate
from caqnal_report.config import load_yaml
from caqnal_report.render.deck import render_all


def _hashes(root: Path) -> dict[str, str]:
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(root.rglob("*")) if p.is_file()}


def test_render_all_decks_and_reproducible(cfg, results, tmp_path):
    outs = []
    for i in range(2):
        out = tmp_path / f"run{i}"
        write_intermediate(results, out / "intermediate", cfg)
        res = load_intermediate(out / "intermediate")
        written = render_all(cfg, res, out, cfg.decks + ["excel"], pdf=False)
        outs.append(out)
        assert any(p.suffix == ".xlsx" for p in written)
        deck = load_yaml(cfg.deck_path("全庁"))
        prs = Presentation(str(out / f"{cfg.client['name']}_{cfg.fiscal_year}_全庁分析レポート.pptx"))
        assert len(prs.slides) == len(deck["slides"])
        # 部局別は n<5 の組織を作らない
        dept = list((out / "部局別").glob("*.pptx"))
        assert dept and not any("監査委員事務局" in p.name for p in dept)
        assert all(len(Presentation(str(p)).slides) == 7 for p in dept)
    a, b = _hashes(outs[0]), _hashes(outs[1])
    assert a == b, [k for k in a if a.get(k) != b.get(k)]
